import json
import math
import io
from datetime import timedelta

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from django.db import transaction
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, get_object_or_404
from django.views.decorators.http import require_http_methods
from django.views.decorators.csrf import csrf_exempt, ensure_csrf_cookie

from .models import Pozo, IntervaloRevestimiento, ComponenteSarta
from .models_daily_reports import (
    ReporteDiario, PropiedadExtraFluido, WellSurveyStation, WellFormationTop,
    ReporteDiarioBomba, ReporteDiarioBitData, ReporteDiarioBoquilla,
    ReporteDiarioMudConfig, ReporteDiarioMudCheck, ReporteDiarioMudExtraValue,
    TramoSarta, ReporteDiarioComentarios,
)
from .geometria_pozo import construir_perfil_confinamiento, calcular_geometria


@ensure_csrf_cookie
def daily_reports_hub_view(request, pk):
    """Gateway principal para Drilling Fluids and Equipment (Historial y Setup)."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reportes = pozo.reportes_diarios.all().order_by('-fecha')

    ultimo = reportes.first()
    if ultimo:
        sugerencia_fecha = (ultimo.fecha + timedelta(days=1)).isoformat()
        sugerencia_lodo = ultimo.tipo_lodo
    elif hasattr(pozo, 'well_header_info') and pozo.well_header_info and pozo.well_header_info.spud_date:
        sugerencia_fecha = pozo.well_header_info.spud_date.isoformat()
        sugerencia_lodo = 'WBM'
    else:
        from datetime import date
        sugerencia_fecha = date.today().isoformat()
        sugerencia_lodo = 'WBM'

    spud_date_str = None
    if hasattr(pozo, 'well_header_info') and pozo.well_header_info and pozo.well_header_info.spud_date:
        spud_date_str = pozo.well_header_info.spud_date.isoformat()

    return render(request, 'operaciones/avances_19_sep/daily_reports_hub.html', {
        'pozo': pozo,
        'reportes': reportes,
        'sugerencia_fecha': sugerencia_fecha,
        'sugerencia_lodo': sugerencia_lodo,
        'spud_date_str': spud_date_str,
    })


def obtener_o_crear_pumps_bits(reporte):
    """
    Obtiene los datos de bombas, barrena y boquillas del reporte diario.
    Si no existen, los clona del reporte anterior o inicializa defaults de ONE-TRAX.
    """
    bombas = list(reporte.bombas.all().order_by('numero_bomba'))
    if not bombas:
        prev = ReporteDiario.objects.filter(pozo=reporte.pozo, fecha__lt=reporte.fecha).order_by('-fecha').first()
        if prev and prev.bombas.exists():
            for pb in prev.bombas.all().order_by('numero_bomba'):
                b = ReporteDiarioBomba.objects.create(
                    reporte=reporte,
                    numero_bomba=pb.numero_bomba,
                    make_model=pb.make_model,
                    liner_diameter=pb.liner_diameter,
                    stroke_length=pb.stroke_length,
                    rod_diam_duplex=pb.rod_diam_duplex,
                    eficiencia_pct=pb.eficiencia_pct,
                    pump_rate_spm=pb.pump_rate_spm,
                    pump_on_report=pb.pump_on_report,
                    riser_pump=pb.riser_pump
                )
                bombas.append(b)
        else:
            default_pumps = [
                (1, "EMSCO F-1000", 6.5, 12.0, 0.0, 97.0, 80.0, True, False),
                (2, "EMSCO F-1000", 6.5, 12.0, 0.0, 97.0, 80.0, True, False),
                (3, "EMSCO F-1000", 6.0, 12.0, 0.0, 97.0, 0.0, False, False),
                (4, "", 6.0, 12.0, 0.0, 97.0, 0.0, False, False),
            ]
            for num, model, liner, stroke, rod, eff, spm, on_rep, riser in default_pumps:
                b = ReporteDiarioBomba.objects.create(
                    reporte=reporte,
                    numero_bomba=num,
                    make_model=model,
                    liner_diameter=liner,
                    stroke_length=stroke,
                    rod_diam_duplex=rod,
                    eficiencia_pct=eff,
                    pump_rate_spm=spm,
                    pump_on_report=on_rep,
                    riser_pump=riser
                )
                bombas.append(b)

    bit_data = getattr(reporte, 'bit_data', None)
    if not bit_data:
        prev = ReporteDiario.objects.filter(pozo=reporte.pozo, fecha__lt=reporte.fecha).order_by('-fecha').first()
        if prev and hasattr(prev, 'bit_data') and prev.bit_data:
            pb = prev.bit_data
            bit_data = ReporteDiarioBitData.objects.create(
                reporte=reporte,
                bit_description=pb.bit_description,
                bit_number=pb.bit_number,
                washout_pct=pb.washout_pct,
                bit_size=pb.bit_size,
                washout_hole_size=pb.washout_hole_size,
                bit_serial_no=pb.bit_serial_no,
                bit_iadc_code=pb.bit_iadc_code,
                bit_manufacturer=pb.bit_manufacturer,
                rotary_rpm=pb.rotary_rpm,
                rotating_hours=pb.rotating_hours,
                weight_on_bit=pb.weight_on_bit,
                rop=pb.rop,
                riser_pump_flow_rate=pb.riser_pump_flow_rate,
                pump_flow_rate=pb.pump_flow_rate,
                pump_pressure=pb.pump_pressure,
                dp_mwd=pb.dp_mwd,
                dp_motor=pb.dp_motor,
                motor_rpm=pb.motor_rpm,
                surface_code=pb.surface_code,
                surface_pressure=pb.surface_pressure,
                ref_flow_rate=pb.ref_flow_rate,
                on_off_bottom_pressure=pb.on_off_bottom_pressure,
            )
        else:
            total_gpm = sum(b.caudal_gpm for b in bombas)
            bit_data = ReporteDiarioBitData.objects.create(
                reporte=reporte,
                bit_description="Hycalog X-175",
                bit_number="3",
                washout_pct=5.0,
                bit_size=12.25,
                washout_hole_size=12.553,
                rotary_rpm=150.0,
                rotating_hours=0.0,
                weight_on_bit=24000.0,
                rop=30.0,
                pump_flow_rate=round(total_gpm, 1) if total_gpm else 803.0,
                pump_pressure=3000.0,
                surface_code="1"
            )

    boquillas = list(reporte.boquillas.all().order_by('-size_32nds', 'id'))
    if not boquillas:
        prev = ReporteDiario.objects.filter(pozo=reporte.pozo, fecha__lt=reporte.fecha).order_by('-fecha').first()
        if prev and prev.boquillas.exists():
            for pnoz in prev.boquillas.all():
                noz = ReporteDiarioBoquilla.objects.create(
                    reporte=reporte,
                    size_32nds=pnoz.size_32nds,
                    cantidad=pnoz.cantidad
                )
                boquillas.append(noz)
        else:
            n1 = ReporteDiarioBoquilla.objects.create(reporte=reporte, size_32nds=14, cantidad=2)
            n2 = ReporteDiarioBoquilla.objects.create(reporte=reporte, size_32nds=13, cantidad=3)
            boquillas = [n1, n2]

    return bombas, bit_data, boquillas


def obtener_o_crear_mud_checks(reporte):
    """
    Obtiene la configuración de lodo y los 4 chequeos diarios (Tab #3).
    Si no existen, los clona del reporte anterior o los inicializa con defaults de ONE-TRAX.
    """
    mud_config, _ = ReporteDiarioMudConfig.objects.get_or_create(reporte=reporte)
    checks = list(reporte.mud_checks.all().order_by('check_number'))

    if len(checks) < 4:
        existing_numbers = {c.check_number for c in checks}
        prev = ReporteDiario.objects.filter(pozo=reporte.pozo, fecha__lt=reporte.fecha).order_by('-fecha').first()
        
        defaults_by_num = {
            1: {
                'is_primary': True,
                'sample_from': 'In',
                'time_taken': '09:00',
                'flowline_temp': 90.0,
                'depth': reporte.profundidad_actual if reporte.profundidad_actual > 0 else 5002.0,
                'tvd': reporte.profundidad_tvd if reporte.profundidad_tvd > 0 else 4750.0,
                'mud_weight': 10.5,
                'mw_temp': 86.0,
                'funnel_viscosity': 45.0,
                'rheology_temp': 120.0,
                'r600': 50.0,
                'r300': 38.0,
                'r200': 26.0,
                'r100': 22.0,
                'r6': 17.0,
                'r3': 14.0,
                'pv': 12.0,
                'yp': 26.0,
                'gel_10s': 14.0,
                'gel_10m': 16.0,
                'gel_30m': 18.0,
                'api_fluid_loss': 6.0,
                'hthp_fluid_loss': 14.0,
                'cake_api': 1.0,
                'cake_hthp': 2.0,
                'solids_pct': 12.0,
                'oil_pct': 0.0,
                'water_pct': 88.0,
                'sand_pct': 0.25,
                'ph': 9.5,
                'ph_temp': 80.0,
                'pm': 1.2,
                'pf': 0.4,
                'mf': 1.8,
                'chlorides': 12000.0,
                'calcium_hardness': 240.0,
                'mbt': 15.0,
            },
            2: {
                'is_primary': False,
                'sample_from': 'Flowline',
                'time_taken': '15:00',
                'flowline_temp': 95.0,
                'depth': reporte.profundidad_actual if reporte.profundidad_actual > 0 else 5002.0,
                'tvd': reporte.profundidad_tvd if reporte.profundidad_tvd > 0 else 4750.0,
                'mud_weight': 10.5,
                'funnel_viscosity': 45.0,
                'rheology_temp': 120.0,
            },
            3: {
                'is_primary': False,
                'sample_from': 'Suction',
                'time_taken': '21:00',
                'flowline_temp': 88.0,
                'depth': reporte.profundidad_actual if reporte.profundidad_actual > 0 else 5002.0,
                'tvd': reporte.profundidad_tvd if reporte.profundidad_tvd > 0 else 4750.0,
                'mud_weight': 10.5,
                'funnel_viscosity': 45.0,
                'rheology_temp': 120.0,
            },
            4: {
                'is_primary': False,
                'sample_from': '',
                'time_taken': '03:00',
                'depth': reporte.profundidad_actual if reporte.profundidad_actual > 0 else 5002.0,
                'tvd': reporte.profundidad_tvd if reporte.profundidad_tvd > 0 else 4750.0,
                'mud_weight': 10.5,
                'funnel_viscosity': 45.0,
                'rheology_temp': 120.0,
            }
        }

        for num in range(1, 5):
            if num not in existing_numbers:
                vals = defaults_by_num[num]
                if prev and prev.mud_checks.filter(check_number=num).exists():
                    p_chk = prev.mud_checks.get(check_number=num)
                    vals = {
                        'is_primary': p_chk.is_primary,
                        'sample_from': p_chk.sample_from,
                        'time_taken': p_chk.time_taken,
                        'flowline_temp': p_chk.flowline_temp,
                        'depth': reporte.profundidad_actual or p_chk.depth,
                        'tvd': reporte.profundidad_tvd or p_chk.tvd,
                        'mud_weight': p_chk.mud_weight,
                        'mw_temp': p_chk.mw_temp,
                        'funnel_viscosity': p_chk.funnel_viscosity,
                        'rheology_temp': p_chk.rheology_temp,
                        'r600': p_chk.r600,
                        'r300': p_chk.r300,
                        'r200': p_chk.r200,
                        'r100': p_chk.r100,
                        'r6': p_chk.r6,
                        'r3': p_chk.r3,
                        'pv': p_chk.pv,
                        'yp': p_chk.yp,
                        'gel_10s': p_chk.gel_10s,
                        'gel_10m': p_chk.gel_10m,
                        'gel_30m': p_chk.gel_30m,
                        'api_fluid_loss': p_chk.api_fluid_loss,
                        'hthp_fluid_loss': p_chk.hthp_fluid_loss,
                        'cake_api': p_chk.cake_api,
                        'cake_hthp': p_chk.cake_hthp,
                        'solids_pct': p_chk.solids_pct,
                        'oil_pct': p_chk.oil_pct,
                        'water_pct': p_chk.water_pct,
                        'sand_pct': p_chk.sand_pct,
                        'ph': p_chk.ph,
                        'ph_temp': p_chk.ph_temp,
                        'pm': p_chk.pm,
                        'pf': p_chk.pf,
                        'mf': p_chk.mf,
                        'chlorides': p_chk.chlorides,
                        'calcium_hardness': p_chk.calcium_hardness,
                        'mbt': p_chk.mbt,
                        'electrical_stability': p_chk.electrical_stability,
                        'excess_lime': p_chk.excess_lime,
                    }

                chk = ReporteDiarioMudCheck.objects.create(
                    reporte=reporte,
                    check_number=num,
                    **vals
                )
                checks.append(chk)
        checks.sort(key=lambda c: c.check_number)

    return mud_config, checks


@ensure_csrf_cookie
def reporte_diario_detalle_view(request, pk, reporte_pk):
    """Vista de 8 pestañas (Detalle del Reporte Diario)."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    # Navegación entre reportes (Browsing bar)
    reportes_qs = pozo.reportes_diarios.all().order_by('fecha')
    reporte_anterior = pozo.reportes_diarios.filter(fecha__lt=reporte.fecha).order_by('-fecha').first()
    reporte_siguiente = pozo.reportes_diarios.filter(fecha__gt=reporte.fecha).order_by('fecha').first()
    primer_reporte = reportes_qs.first()
    ultimo_reporte = reportes_qs.last()

    spud_date = None
    if hasattr(pozo, 'well_header_info') and pozo.well_header_info:
        spud_date = pozo.well_header_info.spud_date

    formation_tops = pozo.formation_tops.all().order_by('depth_ft', 'id')

    # Datos para Tab 2: Pumps / Bits
    bombas, bit_data, boquillas = obtener_o_crear_pumps_bits(reporte)
    tfa_calculado = sum(b.area_sq_in for b in boquillas)
    total_flow_rate_calculado = sum(b.caudal_gpm for b in bombas)

    # Datos para Tab 3: Mud Properties
    mud_config, mud_checks = obtener_o_crear_mud_checks(reporte)

    # Propiedades extra de lodo activas para el pozo
    props_extra = pozo.propiedades_extra.all().order_by('numero')

    # Datos para Tab 4: Well Geometry
    intervalos_pozo = pozo.intervalos_revestimiento.all().order_by('numero_intervalo')

    return render(request, 'operaciones/avances_19_sep/reporte_diario_detalle.html', {
        'intervalos_pozo': intervalos_pozo,
        'pozo': pozo,
        'reporte': reporte,
        'reporte_anterior': reporte_anterior,
        'reporte_siguiente': reporte_siguiente,
        'primer_reporte': primer_reporte,
        'ultimo_reporte': ultimo_reporte,
        'todos_reportes': reportes_qs,
        'spud_date': spud_date,
        'formation_tops': formation_tops,
        'bombas': bombas,
        'bit_data': bit_data,
        'boquillas': boquillas,
        'tfa_calculado': round(tfa_calculado, 2),
        'total_flow_rate_calculado': round(total_flow_rate_calculado, 1),
        'mud_config': mud_config,
        'mud_checks': mud_checks,
        'props_extra': props_extra,
    })


# =====================================================================
# API: Propiedades Extra de Fluidos (Extra Report Labels)
# =====================================================================

@require_http_methods(["GET"])
def api_propiedades_extra_list(request, pk):
    """Devuelve las propiedades extra configuradas para el pozo en formato de diccionario clave-valor."""
    pozo = get_object_or_404(Pozo, pk=pk)
    props = pozo.propiedades_extra.all()

    data = {}
    for p in props:
        key = f"{p.tipo_fluido}_{p.numero}"
        data[key] = {
            'id': p.id,
            'tipo_fluido': p.tipo_fluido,
            'numero': p.numero,
            'etiqueta': p.etiqueta,
            'unidad': p.unidad,
        }

    return JsonResponse({'ok': True, 'propiedades': data})


@csrf_exempt
@require_http_methods(["POST"])
def api_propiedades_extra_guardar(request, pk):
    """Guarda/actualiza todas las propiedades extra del pozo de una sola vez."""
    pozo = get_object_or_404(Pozo, pk=pk)

    try:
        body = json.loads(request.body)
        items = body.get('items', [])
    except (json.JSONDecodeError, KeyError):
        return JsonResponse({'ok': False, 'error': 'JSON invalido'}, status=400)

    guardados = 0
    for item in items:
        tipo = item.get('tipo_fluido', '')
        numero = item.get('numero', 0)
        etiqueta = item.get('etiqueta', '').strip()
        unidad = item.get('unidad', '').strip()

        if not tipo or not numero:
            continue

        if etiqueta or unidad:
            PropiedadExtraFluido.objects.update_or_create(
                pozo=pozo,
                tipo_fluido=tipo,
                numero=numero,
                defaults={'etiqueta': etiqueta, 'unidad': unidad},
            )
            guardados += 1
        else:
            PropiedadExtraFluido.objects.filter(
                pozo=pozo, tipo_fluido=tipo, numero=numero
            ).delete()

    return JsonResponse({'ok': True, 'guardados': guardados})


# =====================================================================
# API: Reportes Diarios (CRUD)
# =====================================================================

@require_http_methods(["GET"])
def api_reportes_diarios_list(request, pk):
    pozo = get_object_or_404(Pozo, pk=pk)
    reportes = pozo.reportes_diarios.all()

    data = []
    for r in reportes:
        data.append({
            'id': r.id,
            'fecha': r.fecha.isoformat(),
            'tipo_lodo': r.tipo_lodo,
            'tipo_lodo_display': r.get_tipo_lodo_display(),
            'profundidad_actual': r.profundidad_actual,
            'actividad_actual': r.actividad_actual,
        })

    ultimo = reportes.first()
    if ultimo:
        fecha_sug = (ultimo.fecha + timedelta(days=1)).isoformat()
        lodo_sug = ultimo.tipo_lodo
    elif hasattr(pozo, 'well_header_info') and pozo.well_header_info and pozo.well_header_info.spud_date:
        fecha_sug = pozo.well_header_info.spud_date.isoformat()
        lodo_sug = 'WBM'
    else:
        from datetime import date
        fecha_sug = date.today().isoformat()
        lodo_sug = 'WBM'

    sugerencia = {
        'fecha': fecha_sug,
        'tipo_lodo': lodo_sug,
    }
    return JsonResponse({'ok': True, 'reportes': data, 'sugerencia': sugerencia})


@csrf_exempt
@require_http_methods(["POST"])
def api_reporte_diario_crear(request, pk):
    pozo = get_object_or_404(Pozo, pk=pk)
    try:
        body = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({'ok': False, 'error': 'JSON invalido'}, status=400)

    fecha = body.get('fecha')
    tipo_lodo = body.get('tipo_lodo', 'WBM')
    copy_data = body.get('copy_data', True)

    if not fecha:
        return JsonResponse({'ok': False, 'error': 'La fecha es obligatoria'}, status=400)
    if ReporteDiario.objects.filter(pozo=pozo, fecha=fecha).exists():
        return JsonResponse({'ok': False, 'error': f'Ya existe un reporte para {fecha}'}, status=400)

    # Defaults de WellHeaderInfo
    op_rep = ''
    co_rep = ''
    m1_rep = ''
    m2_rep = ''
    if hasattr(pozo, 'well_header_info'):
        h = pozo.well_header_info
        op_rep = h.ingeniero_proyecto or ''
        co_rep = h.contratista or ''
        m1_rep = h.ingeniero_miswaco_1 or ''
        m2_rep = h.ingeniero_miswaco_2 or ''

    reporte_anterior = None
    if copy_data:
        reporte_anterior = ReporteDiario.objects.filter(pozo=pozo, fecha__lt=fecha).first()

    reporte = ReporteDiario(pozo=pozo, fecha=fecha, tipo_lodo=tipo_lodo)

    if reporte_anterior:
        reporte.profundidad_actual = reporte_anterior.profundidad_actual
        reporte.profundidad_tvd = reporte_anterior.profundidad_tvd
        reporte.bit_depth = reporte_anterior.bit_depth
        reporte.actividad_actual = reporte_anterior.actividad_actual
        reporte.tipo_fluido_display = reporte_anterior.tipo_fluido_display
        reporte.litologia = reporte_anterior.litologia
        reporte.operador_representante = reporte_anterior.operador_representante
        reporte.contratista_representante = reporte_anterior.contratista_representante
        reporte.mi_representante_1 = reporte_anterior.mi_representante_1
        reporte.mi_representante_2 = reporte_anterior.mi_representante_2
        reporte.telefono_taladro = reporte_anterior.telefono_taladro
        reporte.telefono_almacen = reporte_anterior.telefono_almacen
        reporte.telefonos = reporte_anterior.telefonos
        reporte.fax_numbers = reporte_anterior.fax_numbers
    else:
        reporte.operador_representante = op_rep
        reporte.contratista_representante = co_rep
        reporte.mi_representante_1 = m1_rep
        reporte.mi_representante_2 = m2_rep

    reporte.save()
    fecha_str = reporte.fecha.isoformat() if hasattr(reporte.fecha, 'isoformat') else str(reporte.fecha)
    return JsonResponse({'ok': True, 'id': reporte.id, 'fecha': fecha_str})


@csrf_exempt
@require_http_methods(["DELETE"])
def api_reporte_diario_eliminar(request, pk, reporte_pk):
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)
    reporte.delete()
    return JsonResponse({'ok': True})


# =====================================================================
# API: Tab 1 - General (Guardar)
# =====================================================================
@csrf_exempt
@require_http_methods(["POST"])
def api_reporte_diario_general_guardar(request, pk, reporte_pk):
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)
    try:
        body = json.loads(request.body)
        
        def parse_float(val, default):
            try:
                return float(val) if val not in ('', None) else default
            except (ValueError, TypeError):
                return default

        reporte.fecha = body.get('fecha', reporte.fecha)
        reporte.profundidad_actual = parse_float(body.get('profundidad_actual'), reporte.profundidad_actual)
        reporte.profundidad_tvd = parse_float(body.get('profundidad_tvd'), reporte.profundidad_tvd)
        reporte.bit_depth = parse_float(body.get('bit_depth'), reporte.bit_depth)
        
        reporte.actividad_actual = body.get('actividad_actual', reporte.actividad_actual)
        reporte.tipo_fluido_display = body.get('tipo_fluido_display', reporte.tipo_fluido_display)
        reporte.litologia = body.get('litologia', reporte.litologia)
        reporte.operador_representante = body.get('operador_representante', reporte.operador_representante)
        reporte.contratista_representante = body.get('contratista_representante', reporte.contratista_representante)
        reporte.mi_representante_1 = body.get('mi_representante_1', reporte.mi_representante_1)
        reporte.mi_representante_2 = body.get('mi_representante_2', reporte.mi_representante_2)
        reporte.telefono_taladro = body.get('telefono_taladro', reporte.telefono_taladro)
        reporte.telefono_almacen = body.get('telefono_almacen', reporte.telefono_almacen)
        reporte.telefonos = body.get('telefonos', reporte.telefonos)
        reporte.fax_numbers = body.get('fax_numbers', reporte.fax_numbers)
        reporte.save()
        return JsonResponse({'ok': True})
    except Exception as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=400)


# =====================================================================
# Well Survey: Trayectoria Direccional y Cálculo de TVD
# =====================================================================

def calcular_survey_curvatura_minima(estaciones_raw):
    """
    Calcula TVD, DLS (°/100ft) y Sección Vertical usando el método estándar
    de Curvatura Mínima de la industria petrolera (API).
    """
    estaciones_calculadas = []
    # Ordenar por MD
    sorted_data = sorted(estaciones_raw, key=lambda s: float(s.get('md_ft', 0) or 0))

    tvd_acum = 0.0
    vs_acum = 0.0
    n_s_acum = 0.0
    e_w_acum = 0.0

    prev_md = 0.0
    prev_inc_rad = 0.0
    prev_az_rad = 0.0

    for idx, st in enumerate(sorted_data):
        md = float(st.get('md_ft', 0) or 0)
        inc_deg = float(st.get('inclinacion_deg', 0) or 0)
        az_deg = float(st.get('azimut_deg', 0) or 0)
        comentarios = st.get('comentarios', '')

        inc_rad = math.radians(inc_deg)
        az_rad = math.radians(az_deg)

        if idx == 0:
            if md == 0:
                tvd = 0.0
                dls = 0.0
                vs = 0.0
            else:
                tvd = md * math.cos(inc_rad)
                vs = md * math.sin(inc_rad)
                dls = 0.0
            tvd_acum = tvd
            vs_acum = vs
        else:
            delta_md = md - prev_md
            if delta_md <= 0:
                tvd = tvd_acum
                dls = 0.0
                vs = vs_acum
            else:
                # Curvatura mínima
                cos_beta = math.cos(inc_rad - prev_inc_rad) - math.sin(prev_inc_rad) * math.sin(inc_rad) * (1.0 - math.cos(az_rad - prev_az_rad))
                cos_beta = max(-1.0, min(1.0, cos_beta))
                beta = math.acos(cos_beta)

                if beta > 1e-6:
                    rf = (2.0 / beta) * math.tan(beta / 2.0)
                    dls = (math.degrees(beta) / delta_md) * 100.0
                else:
                    rf = 1.0
                    dls = 0.0

                delta_tvd = (delta_md / 2.0) * (math.cos(prev_inc_rad) + math.cos(inc_rad)) * rf
                delta_ns = (delta_md / 2.0) * (math.sin(prev_inc_rad) * math.cos(prev_az_rad) + math.sin(inc_rad) * math.cos(az_rad)) * rf
                delta_ew = (delta_md / 2.0) * (math.sin(prev_inc_rad) * math.sin(prev_az_rad) + math.sin(inc_rad) * math.sin(az_rad)) * rf

                tvd_acum += delta_tvd
                n_s_acum += delta_ns
                e_w_acum += delta_ew
                tvd = tvd_acum
                vs = math.sqrt(n_s_acum**2 + e_w_acum**2)

        prev_md = md
        prev_inc_rad = inc_rad
        prev_az_rad = az_rad

        estaciones_calculadas.append({
            'estacion': idx + 1,
            'md_ft': round(md, 2),
            'inclinacion_deg': round(inc_deg, 2),
            'azimut_deg': round(az_deg, 2),
            'tvd_ft': round(tvd, 2),
            'dls_deg_100ft': round(dls, 2),
            'seccion_vertical_ft': round(vs, 2),
            'comentarios': comentarios
        })

    return estaciones_calculadas


@require_http_methods(["GET"])
def api_well_survey_list(request, pk):
    """Devuelve las estaciones de survey registradas para el pozo."""
    pozo = get_object_or_404(Pozo, pk=pk)
    stations_qs = pozo.survey_stations.all().order_by('md_ft', 'estacion')

    if not stations_qs.exists():
        # Retornar estación de superficie inicial por defecto
        data = [{
            'id': None,
            'estacion': 1,
            'md_ft': 0.0,
            'inclinacion_deg': 0.0,
            'azimut_deg': 0.0,
            'tvd_ft': 0.0,
            'dls_deg_100ft': 0.0,
            'seccion_vertical_ft': 0.0,
            'comentarios': 'Superficie (Tie-In)'
        }]
    else:
        data = []
        for s in stations_qs:
            data.append({
                'id': s.id,
                'estacion': s.estacion,
                'md_ft': s.md_ft,
                'inclinacion_deg': s.inclinacion_deg,
                'azimut_deg': s.azimut_deg,
                'tvd_ft': s.tvd_ft,
                'dls_deg_100ft': s.dls_deg_100ft,
                'seccion_vertical_ft': s.seccion_vertical_ft,
                'comentarios': s.comentarios,
            })

    return JsonResponse({'ok': True, 'estaciones': data})


@csrf_exempt
@require_http_methods(["POST"])
def api_well_survey_guardar(request, pk):
    """Guarda y recalcula todas las estaciones direccionales del pozo."""
    pozo = get_object_or_404(Pozo, pk=pk)
    try:
        body = json.loads(request.body)
        raw_stations = body.get('estaciones', [])

        if not raw_stations:
            return JsonResponse({'ok': False, 'error': 'No se enviaron estaciones.'}, status=400)

        # Recalcular matemáticamente con curvatura mínima
        calculadas = calcular_survey_curvatura_minima(raw_stations)

        # Persistir en la base de datos
        pozo.survey_stations.all().delete()
        nuevas_instancias = []
        for item in calculadas:
            nuevas_instancias.append(WellSurveyStation(
                pozo=pozo,
                estacion=item['estacion'],
                md_ft=item['md_ft'],
                inclinacion_deg=item['inclinacion_deg'],
                azimut_deg=item['azimut_deg'],
                tvd_ft=item['tvd_ft'],
                dls_deg_100ft=item['dls_deg_100ft'],
                seccion_vertical_ft=item['seccion_vertical_ft'],
                comentarios=item['comentarios']
            ))
        WellSurveyStation.objects.bulk_create(nuevas_instancias)

        # Último TVD calculado
        ultimo_tvd = calculadas[-1]['tvd_ft'] if calculadas else 0.0

        return JsonResponse({
            'ok': True,
            'estaciones': calculadas,
            'ultimo_tvd': ultimo_tvd,
            'mensaje': f"{len(calculadas)} estaciones guardadas y calculadas exitosamente."
        })
    except Exception as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=400)


# =====================================================================
# API: Cost Overview (Pantalla Detallada de Costos)
# =====================================================================

@require_http_methods(["GET"])
def api_cost_overview_detail(request, pk, reporte_pk):
    """Devuelve el balance detallado de costos para el modal de Cost Overview."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    moneda = pozo.moneda_simbolo or "USD"

    # Datos estructurados del desglose
    daily_data = {
        'df_chem': 48430.00,
        'ife_sc': 1450.00,
        'drilling_total': 49880.00,
        'df_equip': 1774.00,
        'other_cost': 7565.06,
        'total': 59219.06,
    }
    cumulative_data = {
        'df_chem': 345735.50,
        'ife_sc': 5825.00,
        'drilling_total': 351560.50,
        'df_equip': 6622.00,
        'other_cost': 10464.29,
        'total': 368646.79,
    }

    desglose_items = [
        {'categoria': 'DF Chemicals', 'descripcion': 'Barite (Sacos 100 lb)', 'cantidad': 120, 'unidad': 'SX', 'costo_unitario': 25.50, 'costo_total': 3060.00},
        {'categoria': 'DF Chemicals', 'descripcion': 'Bentonite Premium Gel', 'cantidad': 40, 'unidad': 'SX', 'costo_unitario': 32.00, 'costo_total': 1280.00},
        {'categoria': 'DF Chemicals', 'descripcion': 'Caustic Soda Flakes', 'cantidad': 15, 'unidad': 'SX', 'costo_unitario': 45.00, 'costo_total': 675.00},
        {'categoria': 'DF Chemicals', 'descripcion': 'VersaClean Base Oil Makeup', 'cantidad': 85, 'unidad': 'BBL', 'costo_unitario': 142.00, 'costo_total': 12070.00},
        {'categoria': 'DF Chemicals', 'descripcion': 'Otros Aditivos y Polímeros', 'cantidad': 1, 'unidad': 'GLO', 'costo_unitario': 30495.00, 'costo_total': 30495.00},
        {'categoria': 'DF Personnel', 'descripcion': 'Primary Mud Engineer Day Rate', 'cantidad': 1, 'unidad': 'DAY', 'costo_unitario': 850.00, 'costo_total': 850.00},
        {'categoria': 'IFE/SC Engineer', 'descripcion': 'Solids Control Specialist', 'cantidad': 1, 'unidad': 'DAY', 'costo_unitario': 1450.00, 'costo_total': 1450.00},
        {'categoria': 'DF Equip. /Screens', 'descripcion': 'Shaker Screens API 140 Mesh', 'cantidad': 4, 'unidad': 'EA', 'costo_unitario': 240.00, 'costo_total': 960.00},
        {'categoria': 'DF Equip. /Screens', 'descripcion': 'Centrifuge CD-500 Rental', 'cantidad': 1, 'unidad': 'DAY', 'costo_unitario': 814.00, 'costo_total': 814.00},
        {'categoria': 'Other Cost (DWM/CF)', 'descripcion': 'Cuttings Haul-off / Disposal Box', 'cantidad': 2, 'unidad': 'BOX', 'costo_unitario': 3782.53, 'costo_total': 7565.06},
    ]

    return JsonResponse({
        'ok': True,
        'pozo': pozo.nombre,
        'reporte_numero': reporte.numero_reporte,
        'fecha': reporte.fecha.strftime('%d/%m/%Y'),
        'moneda': moneda,
        'daily': daily_data,
        'cumulative': cumulative_data,
        'items': desglose_items,
    })


# =====================================================================
# Exportación a Excel (.xlsx): Official M-I SWACO ONE-TRAX Report
# =====================================================================

def reporte_diario_excel_view(request, pk, reporte_pk):
    """Genera y descarga el reporte diario oficial en formato Excel (.xlsx)."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    spud_date_str = "-"
    if hasattr(pozo, 'well_header_info') and pozo.well_header_info and pozo.well_header_info.spud_date:
        spud_date_str = pozo.well_header_info.spud_date.strftime('%d/%m/%Y')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Reporte #{reporte.numero_reporte}"

    # Configuración de estilos
    color_navy = "0F172A"
    color_orange = "F26419"
    color_light_gray = "F8FAFC"
    color_border = "CBD5E1"

    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=10, bold=True, color="334155")
    font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=10, color="0F172A")
    font_cost_num = Font(name="Calibri", size=10, bold=True, color="0F172A")

    fill_title = PatternFill(start_color=color_navy, end_color=color_navy, fill_type="solid")
    fill_section = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_cost_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_total_head = PatternFill(start_color="FFF7ED", end_color="FFF7ED", fill_type="solid")
    fill_row_alt = PatternFill(start_color=color_light_gray, end_color=color_light_gray, fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=color_border),
        right=Side(style='thin', color=color_border),
        top=Side(style='thin', color=color_border),
        bottom=Side(style='thin', color=color_border)
    )

    # 1. Título y Banner Superior
    ws.merge_cells("A1:G1")
    cell_title = ws["A1"]
    cell_title.value = f"M-I SWACO • ONE-TRAX DAILY DRILLING FLUIDS & EQUIPMENT REPORT"
    cell_title.font = font_title
    cell_title.fill = fill_title
    cell_title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # 2. Información del Pozo y Cabecera
    operador_nombre = "M-I SWACO"
    if hasattr(pozo, 'well_header_info') and pozo.well_header_info and pozo.well_header_info.operador:
        operador_nombre = pozo.well_header_info.operador

    ws.merge_cells("A2:C2")
    ws["A2"] = f"Pozo: {pozo.nombre}  |  Operador: {operador_nombre}"
    ws["A2"].font = font_bold
    ws["A2"].alignment = Alignment(vertical="center")

    ws.merge_cells("D2:G2")
    ws["D2"] = f"Fecha: {reporte.fecha.strftime('%d/%m/%Y')}  |  Reporte #: {reporte.numero_reporte}"
    ws["D2"].font = font_bold
    ws["D2"].alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[2].height = 22

    # 3. Sección: Parámetros Operacionales
    ws.merge_cells("A4:G4")
    ws["A4"] = "1. PARÁMETROS DE OPERACIÓN Y GEOMETRÍA"
    ws["A4"].font = font_section
    ws["A4"].fill = fill_section
    ws["A4"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[4].height = 24

    op_fields = [
        ("Fecha del Reporte (Date)", reporte.fecha.strftime('%d/%m/%Y'), "Profundidad Medida (Depth)", f"{reporte.profundidad_actual} ft"),
        ("Fecha de Spud (Spud Date)", spud_date_str, "Profundidad Vertical (TVD)", f"{reporte.profundidad_tvd} ft"),
        ("Número de Reporte (Report #)", f"#{reporte.numero_reporte}", "Profundidad de Mecha (Bit Depth)", f"{reporte.bit_depth} ft"),
        ("Actividad Actual (Activity)", reporte.actividad_actual or "-", "Tipo de Fluido (Fluid Type)", reporte.tipo_fluido_display or "-"),
        ("Litología / Formación", reporte.litologia or "-", "Tipo de Chequeo (Mud Check)", reporte.get_tipo_lodo_display()),
    ]

    current_row = 5
    for lbl1, val1, lbl2, val2 in op_fields:
        ws.cell(row=current_row, column=1, value=lbl1).font = font_header
        ws.cell(row=current_row, column=2, value=val1).font = font_bold
        ws.cell(row=current_row, column=4, value=lbl2).font = font_header
        ws.cell(row=current_row, column=5, value=val2).font = font_bold
        for c in range(1, 8):
            ws.cell(row=current_row, column=c).border = thin_border
        current_row += 1

    # 4. Sección: Representantes y Contactos
    current_row += 1
    ws.merge_cells(f"A{current_row}:G{current_row}")
    ws[f"A{current_row}"] = "2. PERSONAL DE GUARDIA Y CONTACTOS"
    ws[f"A{current_row}"].font = font_section
    ws[f"A{current_row}"].fill = fill_section
    ws[f"A{current_row}"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[current_row].height = 24
    current_row += 1

    rep_fields = [
        ("Representante Operador (Operator Rep)", reporte.operador_representante or "-", "Teléfono del Taladro (Rig Phone)", reporte.telefono_taladro or "-"),
        ("Representante Contratista (Contractor)", reporte.contratista_representante or "-", "Teléfono de Almacén (Whse Phone)", reporte.telefono_almacen or "-"),
        ("Ing. M-I SWACO 1 (Primary Mud Eng)", reporte.mi_representante_1 or "-", "Teléfonos Generales (Telephones)", reporte.telefonos or "-"),
        ("Ing. M-I SWACO 2 (Night Mud Eng)", reporte.mi_representante_2 or "-", "Pagers / Números FAX", reporte.fax_numbers or "-"),
    ]

    for lbl1, val1, lbl2, val2 in rep_fields:
        ws.cell(row=current_row, column=1, value=lbl1).font = font_header
        ws.cell(row=current_row, column=2, value=val1).font = font_regular
        ws.cell(row=current_row, column=4, value=lbl2).font = font_header
        ws.cell(row=current_row, column=5, value=val2).font = font_regular
        for c in range(1, 8):
            ws.cell(row=current_row, column=c).border = thin_border
        current_row += 1

    # 5. Sección: Balance de Costos (Resumen Común)
    current_row += 1
    ws.merge_cells(f"A{current_row}:G{current_row}")
    ws[f"A{current_row}"] = f"3. BALANCE ECONÓMICO DE FLUIDOS Y EQUIPOS ({pozo.moneda_simbolo or 'USD'})"
    ws[f"A{current_row}"].font = font_section
    ws[f"A{current_row}"].fill = fill_section
    ws[f"A{current_row}"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[current_row].height = 24
    current_row += 1

    cost_headers = [
        "Concepto",
        "DF Chem /Personnel",
        "IFE/SC Engineer",
        "Drilling Total",
        "DF Equip. /Screens",
        "Other Cost (DWM/CF)",
        "TOTAL"
    ]
    for col_idx, h in enumerate(cost_headers, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_total_head if col_idx == 7 else fill_cost_header
        cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[current_row].height = 22
    current_row += 1

    cost_rows = [
        ("Daily Cost", 48430.00, 1450.00, 49880.00, 1774.00, 7565.06, 59219.06),
        ("Cumulative Cost", 345735.50, 5825.00, 351560.50, 6622.00, 10464.29, 368646.79),
    ]

    for label, c1, c2, c3, c4, c5, c_tot in cost_rows:
        ws.cell(row=current_row, column=1, value=label).font = font_bold
        ws.cell(row=current_row, column=1).border = thin_border

        for c_idx, val in enumerate([c1, c2, c3, c4, c5, c_tot], start=2):
            cell = ws.cell(row=current_row, column=c_idx, value=val)
            cell.font = font_cost_num
            cell.number_format = "$#,##0.00"
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.border = thin_border
            if c_idx == 7:
                cell.fill = fill_total_head
        current_row += 1

    # Ajuste automático de ancho de columnas
    col_widths = {1: 32, 2: 24, 3: 20, 4: 26, 5: 22, 6: 22, 7: 20}
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # Salida en memoria
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Reporte_Diario_{pozo.nombre}_{reporte.fecha.strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# =====================================================================
# API: Lithology Setup (Topes de Formación y Litología)
# =====================================================================

@require_http_methods(["GET"])
def api_formation_tops_list(request, pk):
    """Devuelve los topes de formación y litología configurados para el pozo."""
    pozo = get_object_or_404(Pozo, pk=pk)
    tops_qs = pozo.formation_tops.all().order_by('depth_ft', 'id')
    data = []
    for t in tops_qs:
        data.append({
            'id': t.id,
            'depth_ft': t.depth_ft,
            'formation_top': t.formation_top,
            'lithology': t.lithology,
            'porcentaje_arena': t.porcentaje_arena,
        })
    return JsonResponse({'ok': True, 'formation_tops': data})


@csrf_exempt
@require_http_methods(["POST"])
def api_formation_tops_guardar(request, pk):
    """Guarda/reemplaza la lista completa de topes de formación para el pozo."""
    pozo = get_object_or_404(Pozo, pk=pk)
    try:
        body = json.loads(request.body)
        raw_tops = body.get('formation_tops', [])

        pozo.formation_tops.all().delete()
        nuevos = []
        for item in raw_tops:
            top_name = str(item.get('formation_top', '')).strip()
            if not top_name and not item.get('depth_ft'):
                continue
            depth = float(item.get('depth_ft', 0) or 0)
            lith = str(item.get('lithology', '')).strip()
            sand = float(item.get('porcentaje_arena', 0) or 0)
            nuevos.append(WellFormationTop(
                pozo=pozo,
                depth_ft=depth,
                formation_top=top_name,
                lithology=lith,
                porcentaje_arena=sand
            ))
        nuevos.sort(key=lambda x: x.depth_ft)
        WellFormationTop.objects.bulk_create(nuevos)

        guardados = [{
            'id': t.id,
            'depth_ft': t.depth_ft,
            'formation_top': t.formation_top,
            'lithology': t.lithology,
            'porcentaje_arena': t.porcentaje_arena,
        } for t in pozo.formation_tops.all().order_by('depth_ft', 'id')]

        return JsonResponse({
            'ok': True,
            'formation_tops': guardados,
            'mensaje': f"{len(guardados)} topes de formación guardados exitosamente."
        })
    except Exception as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=400)


# =====================================================================
# API: Tab 2 - Pumps / Bits (Bombas, Barrena, Boquillas y Presiones)
# =====================================================================

@require_http_methods(["GET"])
def api_pumps_bits_detail(request, pk, reporte_pk):
    """Devuelve los datos completos de bombas, barrena, boquillas y parámetros de Tab 2."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    bombas, bit_data, boquillas = obtener_o_crear_pumps_bits(reporte)

    bombas_data = []
    for b in bombas:
        bombas_data.append({
            'id': b.id,
            'numero_bomba': b.numero_bomba,
            'make_model': b.make_model,
            'liner_diameter': b.liner_diameter,
            'stroke_length': b.stroke_length,
            'rod_diam_duplex': b.rod_diam_duplex or 0.0,
            'eficiencia_pct': b.eficiencia_pct,
            'pump_rate_spm': b.pump_rate_spm,
            'pump_on_report': b.pump_on_report,
            'riser_pump': b.riser_pump,
            'desplazamiento_bbl_stk': round(b.desplazamiento_bbl_stk, 5),
            'desplazamiento_gal_stk': round(b.desplazamiento_gal_stk, 4),
            'caudal_gpm': round(b.caudal_gpm, 1),
        })

    boquillas_data = []
    tfa_total = 0.0
    for noz in boquillas:
        area = noz.area_sq_in
        tfa_total += area
        boquillas_data.append({
            'id': noz.id,
            'size_32nds': noz.size_32nds,
            'cantidad': noz.cantidad,
            'area_sq_in': round(area, 4),
        })

    bit_dict = {
        'bit_description': bit_data.bit_description,
        'bit_number': bit_data.bit_number,
        'washout_pct': bit_data.washout_pct,
        'bit_size': bit_data.bit_size,
        'washout_hole_size': bit_data.washout_hole_size,
        'bit_serial_no': bit_data.bit_serial_no,
        'bit_iadc_code': bit_data.bit_iadc_code,
        'bit_manufacturer': bit_data.bit_manufacturer,
        'rotary_rpm': bit_data.rotary_rpm,
        'rotating_hours': bit_data.rotating_hours,
        'weight_on_bit': bit_data.weight_on_bit,
        'rop': bit_data.rop,
        'riser_pump_flow_rate': bit_data.riser_pump_flow_rate,
        'pump_flow_rate': bit_data.pump_flow_rate,
        'pump_pressure': bit_data.pump_pressure,
        'dp_mwd': bit_data.dp_mwd,
        'dp_motor': bit_data.dp_motor,
        'motor_rpm': bit_data.motor_rpm,
        'surface_code': bit_data.surface_code,
        'surface_pressure': bit_data.surface_pressure,
        'ref_flow_rate': bit_data.ref_flow_rate,
        'on_off_bottom_pressure': bit_data.on_off_bottom_pressure,
    }

    total_flow_rate = sum(b['caudal_gpm'] for b in bombas_data)

    return JsonResponse({
        'ok': True,
        'bombas': bombas_data,
        'boquillas': boquillas_data,
        'tfa': round(tfa_total, 2),
        'total_flow_rate_calculado': round(total_flow_rate, 1),
        'bit_data': bit_dict,
    })


@csrf_exempt
@require_http_methods(["POST"])
def api_pumps_bits_guardar(request, pk, reporte_pk):
    """Guarda todos los cambios de bombas, boquillas y barrena del reporte diario."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    try:
        body = json.loads(request.body)

        def parse_flt(val, default=0.0):
            try:
                return float(val) if val not in ('', None) else default
            except (ValueError, TypeError):
                return default

        # 1. Guardar Bombas
        raw_bombas = body.get('bombas', [])
        for b_item in raw_bombas:
            num = int(b_item.get('numero_bomba', 1))
            b_obj, _ = ReporteDiarioBomba.objects.get_or_create(reporte=reporte, numero_bomba=num)
            b_obj.make_model = str(b_item.get('make_model', '')).strip()
            b_obj.liner_diameter = parse_flt(b_item.get('liner_diameter'), 6.5)
            b_obj.stroke_length = parse_flt(b_item.get('stroke_length'), 12.0)
            b_obj.rod_diam_duplex = parse_flt(b_item.get('rod_diam_duplex'), 0.0)
            b_obj.eficiencia_pct = parse_flt(b_item.get('eficiencia_pct'), 97.0)
            b_obj.pump_rate_spm = parse_flt(b_item.get('pump_rate_spm'), 0.0)
            b_obj.pump_on_report = bool(b_item.get('pump_on_report', False))
            b_obj.riser_pump = bool(b_item.get('riser_pump', False))
            b_obj.save()

        # 2. Guardar Boquillas (Reemplazo)
        raw_boquillas = body.get('boquillas', [])
        reporte.boquillas.all().delete()
        nuevas_boquillas = []
        for noz_item in raw_boquillas:
            sz = int(noz_item.get('size_32nds', 0) or 0)
            qty = int(noz_item.get('cantidad', 0) or 0)
            if sz > 0 and qty > 0:
                nuevas_boquillas.append(ReporteDiarioBoquilla(
                    reporte=reporte,
                    size_32nds=sz,
                    cantidad=qty
                ))
        ReporteDiarioBoquilla.objects.bulk_create(nuevas_boquillas)

        # 3. Guardar Bit Data
        raw_bit = body.get('bit_data', {})
        bit_obj, _ = ReporteDiarioBitData.objects.get_or_create(reporte=reporte)

        bit_obj.bit_description = str(raw_bit.get('bit_description', bit_obj.bit_description)).strip()
        bit_obj.bit_number = str(raw_bit.get('bit_number', bit_obj.bit_number)).strip()
        bit_obj.washout_pct = parse_flt(raw_bit.get('washout_pct'), bit_obj.washout_pct)
        bit_obj.bit_size = parse_flt(raw_bit.get('bit_size'), bit_obj.bit_size)

        washout_calc = bit_obj.bit_size * math.sqrt(1.0 + (bit_obj.washout_pct / 100.0))
        bit_obj.washout_hole_size = round(washout_calc, 3)

        bit_obj.bit_serial_no = str(raw_bit.get('bit_serial_no', bit_obj.bit_serial_no)).strip()
        bit_obj.bit_iadc_code = str(raw_bit.get('bit_iadc_code', bit_obj.bit_iadc_code)).strip()
        bit_obj.bit_manufacturer = str(raw_bit.get('bit_manufacturer', bit_obj.bit_manufacturer)).strip()

        bit_obj.rotary_rpm = parse_flt(raw_bit.get('rotary_rpm'), bit_obj.rotary_rpm)
        bit_obj.rotating_hours = parse_flt(raw_bit.get('rotating_hours'), bit_obj.rotating_hours)
        bit_obj.weight_on_bit = parse_flt(raw_bit.get('weight_on_bit'), bit_obj.weight_on_bit)
        bit_obj.rop = parse_flt(raw_bit.get('rop'), bit_obj.rop)
        bit_obj.riser_pump_flow_rate = parse_flt(raw_bit.get('riser_pump_flow_rate'), bit_obj.riser_pump_flow_rate)

        bit_obj.pump_flow_rate = parse_flt(raw_bit.get('pump_flow_rate'), bit_obj.pump_flow_rate)
        bit_obj.pump_pressure = parse_flt(raw_bit.get('pump_pressure'), bit_obj.pump_pressure)
        bit_obj.dp_mwd = parse_flt(raw_bit.get('dp_mwd'), bit_obj.dp_mwd)
        bit_obj.dp_motor = parse_flt(raw_bit.get('dp_motor'), bit_obj.dp_motor)
        bit_obj.motor_rpm = parse_flt(raw_bit.get('motor_rpm'), bit_obj.motor_rpm)

        bit_obj.surface_code = str(raw_bit.get('surface_code', bit_obj.surface_code)).strip()
        bit_obj.surface_pressure = parse_flt(raw_bit.get('surface_pressure'), bit_obj.surface_pressure)
        bit_obj.ref_flow_rate = parse_flt(raw_bit.get('ref_flow_rate'), bit_obj.ref_flow_rate)
        bit_obj.on_off_bottom_pressure = parse_flt(raw_bit.get('on_off_bottom_pressure'), bit_obj.on_off_bottom_pressure)
        bit_obj.save()

        # Recalcular valores
        bombas_actuales = reporte.bombas.all().order_by('numero_bomba')
        total_gpm = sum(b.caudal_gpm for b in bombas_actuales)
        tfa_actual = sum(b.area_sq_in for b in reporte.boquillas.all())

        return JsonResponse({
            'ok': True,
            'mensaje': 'Datos de Bombas e Hidráulica guardados correctamente.',
            'total_flow_rate_calculado': round(total_gpm, 1),
            'washout_hole_size': bit_obj.washout_hole_size,
            'tfa': round(tfa_actual, 2),
        })
    except Exception as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=400)


# =====================================================================
# API: Pestaña 3 - Mud Properties (ONE-TRAX Daily Data Input)
# =====================================================================

@require_http_methods(["GET"])
def api_mud_properties_detail(request, pk, reporte_pk):
    """Devuelve la configuración y los 4 chequeos de lodo del reporte diario."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    mud_config, checks = obtener_o_crear_mud_checks(reporte)

    # Serializar configuración
    config_dict = {
        'solids_equation': mud_config.solids_equation,
        'is_weighted': mud_config.is_weighted,
        'sg_base_oil': mud_config.sg_base_oil,
        'sg_weight_material': mud_config.sg_weight_material,
        'sg_drill_solids': mud_config.sg_drill_solids,
        'obm_salt_type': getattr(mud_config, 'obm_salt_type', 'CaCl2'),
        'tipo_lodo': reporte.tipo_lodo,
    }

    # Serializar chequeos
    checks_list = []
    for c in checks:
        checks_list.append({
            'id': c.id,
            'check_number': c.check_number,
            'is_primary': c.is_primary,
            'sample_from': c.sample_from,
            'time_taken': c.time_taken,
            'flowline_temp': c.flowline_temp,
            'depth': c.depth,
            'tvd': c.tvd,
            'mud_weight': c.mud_weight,
            'mw_temp': c.mw_temp,
            'funnel_viscosity': c.funnel_viscosity,
            'rheology_temp': c.rheology_temp,
            'r600': c.r600,
            'r300': c.r300,
            'r200': c.r200,
            'r100': c.r100,
            'r6': c.r6,
            'r3': c.r3,
            'pv': c.pv,
            'yp': c.yp,
            'gel_10s': c.gel_10s,
            'gel_10m': c.gel_10m,
            'gel_30m': c.gel_30m,
            'api_fluid_loss': c.api_fluid_loss,
            'hthp_fluid_loss': c.hthp_fluid_loss,
            'cake_api': c.cake_api,
            'cake_hthp': c.cake_hthp,
            'solids_pct': c.solids_pct,
            'oil_pct': c.oil_pct,
            'water_pct': c.water_pct,
            'sand_pct': c.sand_pct,
            'retort_mud_weight': c.retort_mud_weight,
            'retort_mud_temp': c.retort_mud_temp,
            'k_from_kcl': c.k_from_kcl,
            'wt_additive_sg': c.wt_additive_sg,
            'oil_sg': c.oil_sg,
            'frac_bent': c.frac_bent,
            'chem_conc': c.chem_conc,
            'drill_solids_sg': c.drill_solids_sg,
            'nacl_pct': c.nacl_pct,
            'nacl_ppb': c.nacl_ppb,
            'kcl_pct': c.kcl_pct,
            'kcl_ppb': c.kcl_ppb,
            'lgs_pct': c.lgs_pct,
            'lgs_ppb': c.lgs_ppb,
            'bentonite_pct': c.bentonite_pct,
            'bentonite_ppb': c.bentonite_ppb,
            'drill_solids_pct': c.drill_solids_pct,
            'drill_solids_ppb': c.drill_solids_ppb,
            'hgs_pct': c.hgs_pct,
            'hgs_ppb': c.hgs_ppb,
            'salt_pct_wt': c.salt_pct_wt,
            'salt_ppb': c.salt_ppb,
            'adjusted_solids_pct': c.adjusted_solids_pct,
            'oil_water_ratio': c.oil_water_ratio,
            'avg_sg_solids': c.avg_sg_solids,
            'ph': c.ph,
            'ph_temp': c.ph_temp,
            'pm': c.pm,
            'pf': c.pf,
            'mf': c.mf,
            'chlorides': c.chlorides,
            'calcium_hardness': c.calcium_hardness,
            'mbt': c.mbt,
            'electrical_stability': c.electrical_stability,
            'excess_lime': c.excess_lime,
        })

    # Propiedades extra y sus valores para cada chequeo
    props_extra = pozo.propiedades_extra.all().order_by('numero')
    extra_props_list = []
    for pe in props_extra:
        vals_by_check = {}
        for c in checks:
            ev = ReporteDiarioMudExtraValue.objects.filter(mud_check=c, propiedad_extra=pe).first()
            vals_by_check[str(c.check_number)] = ev.valor if ev else ""
        extra_props_list.append({
            'id': pe.id,
            'numero': pe.numero,
            'etiqueta': pe.etiqueta,
            'unidad': pe.unidad,
            'orden_impresion': pe.orden_impresion,
            'valores': vals_by_check
        })

    return JsonResponse({
        'ok': True,
        'config': config_dict,
        'checks': checks_list,
        'extra_properties': extra_props_list
    })


@csrf_exempt
@require_http_methods(["POST"])
def api_mud_properties_guardar(request, pk, reporte_pk):
    """Guarda los cambios en configuración, chequeos y propiedades extra de lodo."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    try:
        body = json.loads(request.body)

        def parse_flt(val, default=None):
            if val in ('', None):
                return default
            try:
                return float(val)
            except (ValueError, TypeError):
                return default

        # 1. Guardar Configuración de Lodo y Sólidos
        raw_config = body.get('config', {})
        mud_config, _ = ReporteDiarioMudConfig.objects.get_or_create(reporte=reporte)
        if 'solids_equation' in raw_config:
            mud_config.solids_equation = str(raw_config['solids_equation']).strip()
        if 'is_weighted' in raw_config:
            mud_config.is_weighted = bool(raw_config['is_weighted'])
        if 'sg_base_oil' in raw_config:
            mud_config.sg_base_oil = parse_flt(raw_config['sg_base_oil'], mud_config.sg_base_oil)
        if 'sg_weight_material' in raw_config:
            mud_config.sg_weight_material = parse_flt(raw_config['sg_weight_material'], mud_config.sg_weight_material)
        if 'sg_drill_solids' in raw_config:
            mud_config.sg_drill_solids = parse_flt(raw_config['sg_drill_solids'], mud_config.sg_drill_solids)
        if 'obm_salt_type' in raw_config:
            mud_config.obm_salt_type = str(raw_config['obm_salt_type']).strip()
        mud_config.save()

        # 2. Guardar Chequeos de Lodo (1 a 4)
        raw_checks = body.get('checks', [])
        primary_check_num = int(body.get('primary_check_number', 1) or 1)

        saved_checks_map = {}
        for item in raw_checks:
            num = int(item.get('check_number', 1))
            c_obj, _ = ReporteDiarioMudCheck.objects.get_or_create(reporte=reporte, check_number=num)
            c_obj.is_primary = (num == primary_check_num)

            c_obj.sample_from = str(item.get('sample_from', c_obj.sample_from or '')).strip()
            c_obj.time_taken = str(item.get('time_taken', c_obj.time_taken or '')).strip()
            c_obj.flowline_temp = parse_flt(item.get('flowline_temp'), c_obj.flowline_temp)
            c_obj.depth = parse_flt(item.get('depth'), c_obj.depth)
            c_obj.tvd = parse_flt(item.get('tvd'), c_obj.tvd)

            c_obj.mud_weight = parse_flt(item.get('mud_weight'), c_obj.mud_weight)
            c_obj.mw_temp = parse_flt(item.get('mw_temp'), c_obj.mw_temp)
            c_obj.funnel_viscosity = parse_flt(item.get('funnel_viscosity'), c_obj.funnel_viscosity)

            c_obj.rheology_temp = parse_flt(item.get('rheology_temp'), c_obj.rheology_temp)
            c_obj.r600 = parse_flt(item.get('r600'), c_obj.r600)
            c_obj.r300 = parse_flt(item.get('r300'), c_obj.r300)
            c_obj.r200 = parse_flt(item.get('r200'), c_obj.r200)
            c_obj.r100 = parse_flt(item.get('r100'), c_obj.r100)
            c_obj.r6 = parse_flt(item.get('r6'), c_obj.r6)
            c_obj.r3 = parse_flt(item.get('r3'), c_obj.r3)

            # Autocalcular PV y YP (o aceptar valor explícito)
            if c_obj.r600 is not None and c_obj.r300 is not None:
                c_obj.pv = max(0.0, round(c_obj.r600 - c_obj.r300, 2))
                c_obj.yp = max(0.0, round(c_obj.r300 - c_obj.pv, 2))
            else:
                c_obj.pv = parse_flt(item.get('pv'), c_obj.pv)
                c_obj.yp = parse_flt(item.get('yp'), c_obj.yp)

            c_obj.gel_10s = parse_flt(item.get('gel_10s'), c_obj.gel_10s)
            c_obj.gel_10m = parse_flt(item.get('gel_10m'), c_obj.gel_10m)
            c_obj.gel_30m = parse_flt(item.get('gel_30m'), c_obj.gel_30m)

            c_obj.api_fluid_loss = parse_flt(item.get('api_fluid_loss'), c_obj.api_fluid_loss)
            c_obj.hthp_fluid_loss = parse_flt(item.get('hthp_fluid_loss'), c_obj.hthp_fluid_loss)
            c_obj.cake_api = parse_flt(item.get('cake_api'), c_obj.cake_api)
            c_obj.cake_hthp = parse_flt(item.get('cake_hthp'), c_obj.cake_hthp)

            c_obj.solids_pct = parse_flt(item.get('solids_pct'), c_obj.solids_pct)
            c_obj.oil_pct = parse_flt(item.get('oil_pct'), c_obj.oil_pct)
            c_obj.water_pct = parse_flt(item.get('water_pct'), c_obj.water_pct)
            c_obj.sand_pct = parse_flt(item.get('sand_pct'), c_obj.sand_pct)

            # Entradas de Retorta y Balance de Sólidos (WBM & OBM)
            c_obj.retort_mud_weight = parse_flt(item.get('retort_mud_weight'), c_obj.retort_mud_weight or c_obj.mud_weight)
            c_obj.retort_mud_temp = parse_flt(item.get('retort_mud_temp'), c_obj.retort_mud_temp or 75.0)
            c_obj.k_from_kcl = parse_flt(item.get('k_from_kcl'), c_obj.k_from_kcl or 0.0)
            c_obj.wt_additive_sg = parse_flt(item.get('wt_additive_sg'), c_obj.wt_additive_sg or mud_config.sg_weight_material or 4.20)
            c_obj.oil_sg = parse_flt(item.get('oil_sg'), c_obj.oil_sg or mud_config.sg_base_oil or 0.70)
            c_obj.frac_bent = parse_flt(item.get('frac_bent'), c_obj.frac_bent or 0.1111)
            c_obj.chem_conc = parse_flt(item.get('chem_conc'), c_obj.chem_conc or 0.0)
            c_obj.drill_solids_sg = parse_flt(item.get('drill_solids_sg'), c_obj.drill_solids_sg or mud_config.sg_drill_solids or 2.60)

            c_obj.ph = parse_flt(item.get('ph'), c_obj.ph)
            c_obj.ph_temp = parse_flt(item.get('ph_temp'), c_obj.ph_temp)
            c_obj.pm = parse_flt(item.get('pm'), c_obj.pm)
            c_obj.pf = parse_flt(item.get('pf'), c_obj.pf)
            c_obj.mf = parse_flt(item.get('mf'), c_obj.mf)
            c_obj.chlorides = parse_flt(item.get('chlorides'), c_obj.chlorides)
            c_obj.calcium_hardness = parse_flt(item.get('calcium_hardness'), c_obj.calcium_hardness)
            c_obj.mbt = parse_flt(item.get('mbt'), c_obj.mbt)
            c_obj.electrical_stability = parse_flt(item.get('electrical_stability'), c_obj.electrical_stability)
            c_obj.excess_lime = parse_flt(item.get('excess_lime'), c_obj.excess_lime)

            # Entradas directas OBM si se especifican
            if item.get('salt_pct_wt') is not None: c_obj.salt_pct_wt = parse_flt(item.get('salt_pct_wt'), c_obj.salt_pct_wt)
            if item.get('salt_ppb') is not None: c_obj.salt_ppb = parse_flt(item.get('salt_ppb'), c_obj.salt_ppb)

            # Ejecutar cálculo de balance de sólidos (WBM u OBM según corresponda)
            c_obj.calcular_solids_analysis_automatica()

            # Aceptar sobrescrituras explícitas si vinieron en el payload
            if item.get('nacl_pct') is not None: c_obj.nacl_pct = parse_flt(item.get('nacl_pct'), c_obj.nacl_pct)
            if item.get('nacl_ppb') is not None: c_obj.nacl_ppb = parse_flt(item.get('nacl_ppb'), c_obj.nacl_ppb)
            if item.get('kcl_pct') is not None: c_obj.kcl_pct = parse_flt(item.get('kcl_pct'), c_obj.kcl_pct)
            if item.get('kcl_ppb') is not None: c_obj.kcl_ppb = parse_flt(item.get('kcl_ppb'), c_obj.kcl_ppb)
            if item.get('lgs_pct') is not None: c_obj.lgs_pct = parse_flt(item.get('lgs_pct'), c_obj.lgs_pct)
            if item.get('lgs_ppb') is not None: c_obj.lgs_ppb = parse_flt(item.get('lgs_ppb'), c_obj.lgs_ppb)
            if item.get('bentonite_pct') is not None: c_obj.bentonite_pct = parse_flt(item.get('bentonite_pct'), c_obj.bentonite_pct)
            if item.get('bentonite_ppb') is not None: c_obj.bentonite_ppb = parse_flt(item.get('bentonite_ppb'), c_obj.bentonite_ppb)
            if item.get('drill_solids_pct') is not None: c_obj.drill_solids_pct = parse_flt(item.get('drill_solids_pct'), c_obj.drill_solids_pct)
            if item.get('drill_solids_ppb') is not None: c_obj.drill_solids_ppb = parse_flt(item.get('drill_solids_ppb'), c_obj.drill_solids_ppb)
            if item.get('hgs_pct') is not None: c_obj.hgs_pct = parse_flt(item.get('hgs_pct'), c_obj.hgs_pct)
            if item.get('hgs_ppb') is not None: c_obj.hgs_ppb = parse_flt(item.get('hgs_ppb'), c_obj.hgs_ppb)
            if item.get('adjusted_solids_pct') is not None: c_obj.adjusted_solids_pct = parse_flt(item.get('adjusted_solids_pct'), c_obj.adjusted_solids_pct)
            if item.get('oil_water_ratio') is not None: c_obj.oil_water_ratio = str(item.get('oil_water_ratio')).strip()
            if item.get('avg_sg_solids') is not None: c_obj.avg_sg_solids = parse_flt(item.get('avg_sg_solids'), c_obj.avg_sg_solids)

            c_obj.save()
            saved_checks_map[num] = c_obj

        # 3. Guardar Propiedades Extra si se enviaron
        raw_extra = body.get('extra_values', [])
        for ev_item in raw_extra:
            c_num = int(ev_item.get('check_number', 1))
            prop_id = int(ev_item.get('propiedad_extra_id', 0))
            val_str = str(ev_item.get('valor', '')).strip()

            c_target = saved_checks_map.get(c_num) or ReporteDiarioMudCheck.objects.filter(reporte=reporte, check_number=c_num).first()
            p_extra = PropiedadExtraFluido.objects.filter(id=prop_id, pozo=pozo).first()
            if c_target and p_extra:
                ev_obj, _ = ReporteDiarioMudExtraValue.objects.get_or_create(mud_check=c_target, propiedad_extra=p_extra)
                ev_obj.valor = val_str
                ev_obj.save()

        return JsonResponse({
            'ok': True,
            'mensaje': 'Propiedades de Lodo guardadas exitosamente.',
            'primary_check_number': primary_check_num,
        })
    except Exception as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=400)


# =====================================================================
# API: Pestaña 4 - Well Geometry (ONE-TRAX Daily Data Input)
# =====================================================================

def _float(valor, defecto=0.0):
    try:
        return float(valor) if valor not in ('', None) else defecto
    except (ValueError, TypeError):
        return defecto


def obtener_riser_del_pozo(pozo):
    """
    Devuelve {'longitud_ft', 'id_in'} si el pozo usa riser, o None.

    Si no se capturó la longitud del riser se asume Air Gap + Water Depth, que es
    la distancia real entre la mesa rotaria y el lecho marino.
    """
    header = getattr(pozo, 'well_header_info', None)
    if not header or not getattr(header, 'usa_riser', False):
        return None

    riser_id = _float(header.riser_id_in)
    if riser_id <= 0:
        return None

    longitud = _float(header.riser_length_ft)
    if longitud <= 0:
        longitud = _float(header.air_gap_ft) + _float(header.water_depth_ft)
    if longitud <= 0:
        return None

    return {'longitud_ft': longitud, 'id_in': riser_id}


def obtener_intervalos_del_pozo(pozo):
    """Intervalos de revestimiento del pozo, listos para el motor de geometría."""
    datos = []
    for itv in pozo.intervalos_revestimiento.all().order_by('numero_intervalo'):
        etiqueta = itv.get_tipo_display() if itv.tipo else 'Revestidor'
        if itv.casing_od_in:
            etiqueta = f"{etiqueta} {_float(itv.casing_od_in):g}\""
        datos.append({
            'id': itv.id,
            'numero': itv.numero_intervalo,
            'casing_od_in': _float(itv.casing_od_in),
            'casing_id_in': _float(itv.casing_id_in),
            'hole_size_in': _float(itv.hole_size_in),
            'profundidad_ft': _float(itv.profundidad_ft),
            'tvd_ft': _float(itv.tvd_ft),
            'top_of_liner_ft': _float(itv.top_of_liner_ft),
            'planned_length_ft': _float(itv.planned_length_ft),
            'planned_days': itv.planned_days or 0,
            'interval_days': itv.interval_days or 0,
            'tipo': itv.tipo,
            'etiqueta': etiqueta,
        })
    return datos


def obtener_hidraulica_bombas(reporte):
    """
    Salida de bomba para los cálculos de 'Fondo Arriba' (Bottom Up).

    Las emboladas se cuentan contra UNA bomba (la primera activa del reporte), que es
    como el perforador las cuenta en el taladro; los minutos usan el caudal total de
    todas las bombas activas. Esta combinación reproduce los valores del manual.
    """
    bombas = list(reporte.bombas.all().order_by('numero_bomba'))
    activas = [b for b in bombas if b.pump_on_report]
    bbl_stk = activas[0].desplazamiento_bbl_stk if activas else 0.0
    caudal_gpm = sum(b.caudal_gpm for b in activas)
    return bbl_stk, caudal_gpm


def calcular_geometria_reporte(pozo, reporte):
    """Arma el perfil del pozo y corre el motor de volúmenes para un reporte diario."""
    bit_data = getattr(reporte, 'bit_data', None)

    # Diámetro del hoyo abierto: bit size corregido por washout cuando está disponible.
    hole_size = 0.0
    if bit_data:
        hole_size = _float(bit_data.washout_hole_size) or _float(bit_data.bit_size)
    if hole_size <= 0:
        intervalo = reporte.intervalo_costo
        if intervalo:
            hole_size = _float(intervalo.hole_size_in)

    bit_depth = _float(reporte.bit_depth)
    profundidad = _float(reporte.profundidad_actual)
    pilot_depth = _float(reporte.pilot_hole_depth_ft)
    pilot_size = _float(reporte.pilot_hole_size_in)

    fondo_hoyo = max(profundidad, bit_depth, pilot_depth)

    pilot_hole = None
    if pilot_size > 0 and pilot_depth > profundidad:
        pilot_hole = {
            'hole_size_in': pilot_size,
            'depth_ft': pilot_depth,
            'desde_ft': profundidad,
        }

    riser = obtener_riser_del_pozo(pozo)
    intervalos = obtener_intervalos_del_pozo(pozo)

    perfil = construir_perfil_confinamiento(
        riser=riser,
        intervalos=intervalos,
        fondo_hoyo_ft=fondo_hoyo,
        hole_size_in=hole_size,
        pilot_hole=pilot_hole,
    )

    tramos = [t.to_dict() for t in reporte.tramos_sarta.all().order_by('orden')]
    bbl_stk, caudal_gpm = obtener_hidraulica_bombas(reporte)

    resultado = calcular_geometria(
        perfil_pozo=perfil,
        tramos=tramos,
        bit_depth_ft=bit_depth,
        fondo_hoyo_ft=fondo_hoyo,
        bbl_por_embolada=bbl_stk,
        caudal_gpm=caudal_gpm,
    )
    resultado['perfil_pozo'] = [
        {
            'desde_ft': round(s['desde_ft'], 2),
            'hasta_ft': round(s['hasta_ft'], 2),
            'diametro_in': round(s['diametro_in'], 3),
            'etiqueta': s['etiqueta'],
            'es_hoyo_abierto': s['es_hoyo_abierto'],
        }
        for s in perfil
    ]
    resultado['hole_size_in'] = round(hole_size, 3)
    return resultado, tramos, intervalos


def _resumen_avance_intervalo(reporte, intervalos):
    """Compara lo planeado contra lo real del intervalo de costo asignado al reporte."""
    intervalo = reporte.intervalo_costo
    if not intervalo:
        return None

    datos = next((i for i in intervalos if i['id'] == intervalo.id), None)
    if not datos:
        return None

    planeado_ft = datos['planned_length_ft']
    real_ft = _float(reporte.profundidad_actual)
    porcentaje = round((real_ft / planeado_ft) * 100.0, 1) if planeado_ft > 0 else None

    return {
        'numero': datos['numero'],
        'etiqueta': datos['etiqueta'],
        'casing_od_in': datos['casing_od_in'],
        'casing_id_in': datos['casing_id_in'],
        'hole_size_in': datos['hole_size_in'],
        'profundidad_ft': datos['profundidad_ft'],
        'planned_length_ft': planeado_ft,
        'planned_days': datos['planned_days'],
        'interval_days': datos['interval_days'],
        'avance_ft': real_ft,
        'avance_pct': porcentaje,
    }


@require_http_methods(["GET"])
def api_well_geometry_detail(request, pk, reporte_pk):
    """Devuelve la sarta, el contexto del pozo y los volúmenes calculados del Tab 4."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    geometria, tramos, intervalos = calcular_geometria_reporte(pozo, reporte)

    # Continuidad: de dónde viene el trabajo y cuánto se avanzó hoy.
    anterior = pozo.reportes_diarios.filter(fecha__lt=reporte.fecha).order_by('-fecha').first()
    continuidad = None
    if anterior:
        bit_anterior = _float(anterior.bit_depth)
        continuidad = {
            'fecha_anterior': anterior.fecha.isoformat(),
            'bit_depth_anterior': round(bit_anterior, 2),
            'profundidad_anterior': round(_float(anterior.profundidad_actual), 2),
            'avance_dia_ft': round(_float(reporte.profundidad_actual) - _float(anterior.profundidad_actual), 2),
            'tiene_sarta': anterior.tramos_sarta.exists(),
        }

    bbl_stk, caudal_gpm = obtener_hidraulica_bombas(reporte)

    return JsonResponse({
        'ok': True,
        'tramos': tramos,
        'intervalos': intervalos,
        'riser': obtener_riser_del_pozo(pozo),
        'hole_size_in': geometria['hole_size_in'],
        'bbl_por_embolada': round(bbl_stk, 6),
        'caudal_gpm': round(caudal_gpm, 2),
        'intervalo_costo_id': reporte.intervalo_costo_id,
        'orden_impresion_intervalo': reporte.orden_impresion_intervalo,
        'pilot_hole_size_in': _float(reporte.pilot_hole_size_in),
        'pilot_hole_depth_ft': _float(reporte.pilot_hole_depth_ft),
        'bit_depth_ft': _float(reporte.bit_depth),
        'profundidad_actual_ft': _float(reporte.profundidad_actual),
        'avance_intervalo': _resumen_avance_intervalo(reporte, intervalos),
        'continuidad': continuidad,
        'geometria': geometria,
    })


@csrf_exempt
@require_http_methods(["POST"])
def api_well_geometry_guardar(request, pk, reporte_pk):
    """Guarda la sarta del día (reemplazo total), el intervalo de costo y el hoyo piloto."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    try:
        body = json.loads(request.body)

        # 1. Intervalo de costo y orden de impresión
        intervalo_id = body.get('intervalo_costo_id')
        if intervalo_id in ('', None, 'null'):
            reporte.intervalo_costo = None
        else:
            reporte.intervalo_costo = get_object_or_404(
                IntervaloRevestimiento, pk=int(intervalo_id), pozo=pozo
            )

        orden = body.get('orden_impresion_intervalo')
        reporte.orden_impresion_intervalo = int(orden) if str(orden or '').strip().isdigit() else None

        # 2. Hoyo piloto
        reporte.pilot_hole_size_in = _float(body.get('pilot_hole_size_in'))
        reporte.pilot_hole_depth_ft = _float(body.get('pilot_hole_depth_ft'))
        reporte.save()

        # 3. Sarta: reemplazo total (mismo patrón que boquillas y listas activas)
        filas = body.get('tramos', [])
        nuevos = []
        orden_actual = 0
        for fila in filas:
            longitud = _float(fila.get('longitud_ft'))
            od = _float(fila.get('od_in'))
            descripcion = str(fila.get('descripcion', '')).strip()
            if longitud <= 0 and od <= 0 and not descripcion:
                continue
            orden_actual += 1
            componente_id = fila.get('componente_id')
            nuevos.append(TramoSarta(
                reporte=reporte,
                orden=orden_actual,
                componente_id=int(componente_id) if str(componente_id or '').isdigit() else None,
                descripcion=descripcion[:150],
                longitud_ft=longitud,
                od_in=od,
                id_in=_float(fila.get('id_in')),
                tool_joint_od_in=_float(fila.get('tool_joint_od_in')),
                tool_joint_id_in=_float(fila.get('tool_joint_id_in')),
                tool_joint_length_in=_float(fila.get('tool_joint_length_in')),
                largo_tramo_ft=_float(fila.get('largo_tramo_ft'), 31.0),
            ))

        with transaction.atomic():
            reporte.tramos_sarta.all().delete()
            TramoSarta.objects.bulk_create(nuevos)

        reporte.refresh_from_db()
        geometria, tramos, intervalos = calcular_geometria_reporte(pozo, reporte)

        return JsonResponse({
            'ok': True,
            'mensaje': 'Geometría del pozo guardada correctamente.',
            'tramos': tramos,
            'avance_intervalo': _resumen_avance_intervalo(reporte, intervalos),
            'geometria': geometria,
        })
    except Exception as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=400)


@require_http_methods(["GET"])
def api_sarta_reporte_anterior(request, pk, reporte_pk):
    """Devuelve la sarta del reporte anterior, para heredarla al reporte del día."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    anterior = pozo.reportes_diarios.filter(fecha__lt=reporte.fecha).order_by('-fecha').first()
    if not anterior:
        return JsonResponse({'ok': True, 'tramos': [], 'fecha_origen': None})

    return JsonResponse({
        'ok': True,
        'fecha_origen': anterior.fecha.isoformat(),
        'bit_depth_origen': _float(anterior.bit_depth),
        'tramos': [t.to_dict() for t in anterior.tramos_sarta.all().order_by('orden')],
    })


# =====================================================================
# API: Pestaña 5 - Comments (ONE-TRAX Daily Data Input)
# =====================================================================

def obtener_o_crear_comentarios(reporte):
    """
    Devuelve los comentarios del reporte, creándolos si es la primera vez.

    La especificación de propiedades del lodo (peso, viscosidad, filtrado) es el rango
    objetivo acordado con el operador: no cambia día a día, así que al crear el registro
    se hereda del reporte anterior. Los bloques de texto NO se heredan: cada día tiene
    sus propios comentarios.
    """
    comentarios = getattr(reporte, 'comentarios', None)
    if comentarios:
        return comentarios

    valores = {}
    anterior = (
        ReporteDiario.objects
        .filter(pozo=reporte.pozo, fecha__lt=reporte.fecha)
        .order_by('-fecha')
        .first()
    )
    if anterior:
        previos = getattr(anterior, 'comentarios', None)
        if previos:
            valores = {
                'spec_mud_weight': previos.spec_mud_weight,
                'spec_viscosidad': previos.spec_viscosidad,
                'spec_filtrado': previos.spec_filtrado,
            }

    return ReporteDiarioComentarios.objects.create(reporte=reporte, **valores)


@require_http_methods(["GET"])
def api_comentarios_detail(request, pk, reporte_pk):
    """Devuelve los comentarios del día y el resumen del reporte anterior como referencia."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    comentarios = obtener_o_crear_comentarios(reporte)

    # Resumen del día anterior: sirve de contexto al redactar el de hoy.
    anterior = (
        pozo.reportes_diarios
        .filter(fecha__lt=reporte.fecha)
        .order_by('-fecha')
        .first()
    )
    recap_anterior = None
    if anterior:
        previos = getattr(anterior, 'comentarios', None)
        if previos and previos.mud_recap_remarks.strip():
            recap_anterior = {
                'fecha': anterior.fecha.isoformat(),
                'texto': previos.mud_recap_remarks.strip(),
            }

    return JsonResponse({
        'ok': True,
        'comentarios': comentarios.to_dict(),
        'recap_anterior': recap_anterior,
    })


@csrf_exempt
@require_http_methods(["POST"])
def api_comentarios_guardar(request, pk, reporte_pk):
    """Guarda la especificación de lodo y los tres bloques de comentarios del día."""
    pozo = get_object_or_404(Pozo, pk=pk)
    reporte = get_object_or_404(ReporteDiario, pk=reporte_pk, pozo=pozo)

    try:
        body = json.loads(request.body)
        comentarios = obtener_o_crear_comentarios(reporte)

        comentarios.spec_mud_weight = str(body.get('spec_mud_weight', ''))[:50].strip()
        comentarios.spec_viscosidad = str(body.get('spec_viscosidad', ''))[:50].strip()
        comentarios.spec_filtrado = str(body.get('spec_filtrado', ''))[:50].strip()

        # El Mud Recap se imprime como una sola línea en el Well Recap del pozo:
        # se colapsan los saltos de línea para que no rompa el formato del reporte.
        recap = str(body.get('mud_recap_remarks', '')).strip()
        comentarios.mud_recap_remarks = ' '.join(recap.split())

        comentarios.remarks_and_treatment = str(body.get('remarks_and_treatment', '')).strip()
        comentarios.remarks = str(body.get('remarks', '')).strip()
        comentarios.save()

        return JsonResponse({
            'ok': True,
            'mensaje': 'Comentarios guardados correctamente.',
            'comentarios': comentarios.to_dict(),
        })
    except Exception as e:
        return JsonResponse({'ok': False, 'error': str(e)}, status=400)
