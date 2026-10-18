"""
Excel diario: "Tanques activos" y "Volumen total de circulación" llevan lo calculado de los tanques,
no lo medido (revisión del 15-oct, docs/27_VOLUMEN_CALCULADO_TICKETS_Y_PERDIDAS.md).
El final calculado de un reporte es el inicial del siguiente, aunque ese día se haya medido otro real.
"""

import io
import json
from datetime import date
from decimal import Decimal

import openpyxl
from django.test import TestCase

from operaciones.models import Fosa, Pozo, TipoFosa
from operaciones.models_daily_reports import ReporteDiario
from operaciones.models_inventario import VolumenFosaDia
from operaciones.reporte_excel import HOJA_LODO, generar_reporte_excel

ACTIVA = 1


class ExcelVolumenCalculadoTests(TestCase):

    def setUp(self):
        self.pozo = Pozo.objects.create(nombre='POZO EXCEL VOLUMEN')
        TipoFosa.sembrar_estandar(self.pozo)
        Fosa.objects.create(pozo=self.pozo, numero=1, descripcion='Activo 1', capacidad=Decimal('1000'))
        self.dia1 = self.reporte(date(2026, 10, 15), real='650')
        self.dia2 = self.reporte(date(2026, 10, 16), real='700')
        url = f'/api/pozos/{self.pozo.pk}/daily-report/{self.dia1.pk}/volumetria/transaccion/'
        self.client.post(url, json.dumps({'tipo': 'QUIMICOS', 'fosa': 1, 'agua': 400}),
                         content_type='application/json')

    def reporte(self, fecha, real):
        r = ReporteDiario.objects.create(pozo=self.pozo, fecha=fecha)
        VolumenFosaDia.objects.create(reporte=r, fosa_numero=1, fosa_descripcion='Activo 1',
                                      tipo_codigo=ACTIVA, volumen_final=Decimal(real))
        return r

    def hoja(self, reporte):
        contenido, _ = generar_reporte_excel(self.pozo, reporte)
        return openpyxl.load_workbook(io.BytesIO(contenido))[HOJA_LODO[reporte.tipo_lodo]]

    def test_el_dia_del_movimiento_muestra_lo_calculado_y_no_lo_medido(self):
        ws = self.hoja(self.dia1)
        self.assertEqual(ws['F10'].value, 400)    # medido: 650
        self.assertEqual(ws['E12'].value, 400)

    def test_el_dia_siguiente_parte_del_calculado_del_anterior(self):
        ws = self.hoja(self.dia2)
        self.assertEqual(ws['F10'].value, 400)    # medido: 700
        self.assertEqual(ws['E12'].value, 400)
