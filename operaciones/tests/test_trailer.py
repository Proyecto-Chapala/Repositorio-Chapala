"""
Categoría Trailer (punto 18 de docs/23_CORRECCIONES_REVISION_OCT.md): como el ingeniero, es un
servicio con tarifa diaria y acumulado; no lleva existencias ni volumen, y en el resumen de costos
suma en "Otros costos".
"""

from datetime import date
from decimal import Decimal

from django.test import SimpleTestCase, TestCase

from operaciones import volumetria as vol
from operaciones.models import Pozo, Producto, ProductoActivoPozo
from operaciones.models_daily_reports import ReporteDiario
from operaciones.models_inventario import InventarioProductoDia
from operaciones.recap_pozo import datos_recap
from operaciones.views_inventario import resumen_costos

TRAILER = 7
CATEGORIA_TRAILER = 5


def dia_con_trailer(id_, dias_cobrados, tarifa=250):
    return {'id': id_, 'fecha_texto': f'dia {id_}', 'hoyo_fluido': 0, 'tickets': [], 'transacciones': [],
            'fosas': [{'numero': 1, 'grupo': vol.ACTIVO, 'real': None}],
            'manual': {TRAILER: {'otro': dias_cobrados, 'ajuste': 0, 'precio': tarifa,
                                 'categoria': CATEGORIA_TRAILER}}}


class TrailerTests(SimpleTestCase):

    def test_trailer_es_una_categoria_de_costo(self):
        self.assertEqual(vol.CATEGORIAS_COSTO[CATEGORIA_TRAILER], 'Trailer')

    def test_costo_diario_y_acumulado_sin_existencias(self):
        dias = [dia_con_trailer(1, 1), dia_con_trailer(2, 1)]
        sim = vol.simular(dias, 2, servicios={TRAILER})
        inventario = sim['inventario'][TRAILER]
        self.assertEqual(inventario['costo_diario'], 250)
        self.assertEqual(inventario['costo_acumulado'], 500)
        self.assertEqual(inventario['final'], 0)
        self.assertEqual(sim['costo_dia_cat'][CATEGORIA_TRAILER], 250)
        self.assertEqual(sim['costo_acum_cat'][CATEGORIA_TRAILER], 500)

    def test_no_aporta_volumen(self):
        sim = vol.simular([dia_con_trailer(1, 1)], 1, servicios={TRAILER})
        self.assertEqual(sim['flujos'][vol.ACTIVO].get('quimicos', 0.0), 0.0)

    def test_columnas_del_resumen_de_costos(self):
        self.assertEqual(vol.columna_resumen(1), vol.QUIMICOS_PERSONAL_DF)
        self.assertEqual(vol.columna_resumen(2), vol.QUIMICOS_PERSONAL_DF)
        self.assertEqual(vol.columna_resumen(3), vol.INGENIERIA_IFE_SOLIDOS)
        self.assertEqual(vol.columna_resumen(4), vol.INGENIERIA_IFE_SOLIDOS)
        self.assertEqual(vol.columna_resumen(CATEGORIA_TRAILER), vol.OTROS_COSTOS)
        self.assertEqual(vol.columna_resumen(99), vol.QUIMICOS_PERSONAL_DF)

    def test_todas_las_categorias_tienen_columna(self):
        self.assertEqual(set(vol.CATEGORIAS_COSTO), set(vol.COLUMNA_RESUMEN))


class TrailerEnResumenesTests(TestCase):
    """El trailer cargado en el reporte llega a 'Otros costos' del resumen diario y del recap."""

    def setUp(self):
        self.pozo = Pozo.objects.create(nombre='POZO PRUEBA')
        trailer = Producto.objects.create(codigo='TRAILER', descripcion='Trailer', unidad='DIA')
        ProductoActivoPozo.objects.create(pozo=self.pozo, producto=trailer, unidad='', precio=Decimal('250'),
                                          codigo_costo_diario=CATEGORIA_TRAILER)
        self.reporte = ReporteDiario.objects.create(pozo=self.pozo, fecha=date(2026, 10, 13))
        InventarioProductoDia.objects.create(reporte=self.reporte, producto=trailer, usado_otro=Decimal('1'),
                                             precio=Decimal('250'), categoria_costo=CATEGORIA_TRAILER)

    def test_resumen_de_costos_del_dia(self):
        diario, acumulado, items = resumen_costos(self.pozo, self.reporte)
        self.assertEqual(diario['other_cost'], 250)
        self.assertEqual(diario['df_chem'], 0)
        self.assertEqual(diario['total'], 250)
        self.assertEqual(acumulado['other_cost'], 250)
        self.assertEqual([i['categoria'] for i in items], ['Otros costos'])

    def test_recap_del_pozo(self):
        costos = datos_recap(self.pozo)['costos']
        self.assertEqual(costos['otros'], 250)
        self.assertEqual(costos['total'], 250)

    def test_excel_diario_tiene_fila_de_otros_costos(self):
        import io
        import openpyxl
        from operaciones.reporte_excel import generar_reporte_excel, HOJA_LODO
        contenido, _ = generar_reporte_excel(self.pozo, self.reporte)
        ws = openpyxl.load_workbook(io.BytesIO(contenido))[HOJA_LODO[self.reporte.tipo_lodo]]
        rc = 67   # base agua: primera fila de costos
        self.assertEqual(ws[f'F{rc + 3}'].value, 'Otros (Trailer)')
        self.assertEqual(ws[f'G{rc + 3}'].value, 250)
        self.assertEqual(ws[f'F{rc + 4}'].value, 'TOTALES')
        self.assertEqual(ws[f'G{rc + 4}'].value, 250)
        self.assertEqual(ws[f'I{rc + 4}'].value, 250)
        self.assertIn(f'G{rc + 4}:H{rc + 4}', [str(m) for m in ws.merged_cells.ranges])
