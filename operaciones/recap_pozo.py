"""
Reporte final del pozo (recap / "reporte consolidado").

Lo pidió el ingeniero de AOS: al terminar el pozo (o cuando se quiera, hasta una fecha) se
entrega un reporte que reúne todo lo que pasó: datos del pozo, intervalos, el resumen día
por día, volúmenes y pérdidas, productos consumidos, distribución de tiempo, eventos no
programados, costos, y las conclusiones y recomendaciones del ingeniero.

Este módulo solo ARMA los datos (función datos_recap) y el Excel (generar_recap_excel).
El PDF sale de la página imprimible (views_recap.recap_imprimible_view): el navegador la
guarda como PDF, sin librerías nuevas.

No calcula nada propio: usa los mismos motores que las pestañas (volumetría, mallas,
hoyo perforado), así el recap coincide con los reportes diarios. Las simulaciones se
arman UNA vez para todo el pozo y se consultan día por día en memoria.

Redondeo (pedido del ingeniero): volúmenes sin decimales, densidad con 1 decimal,
costos con 2 decimales (el costo no se redondea más).
"""

import io
from collections import OrderedDict, defaultdict
from datetime import datetime

SECCIONES = OrderedDict([
    ('resumen', 'Datos del pozo y resumen'),
    ('conclusiones', 'Conclusiones y recomendaciones'),
    ('intervalos', 'Resumen por intervalo'),
    ('diario', 'Recap día por día'),
    ('volumenes', 'Volúmenes y pérdidas'),
    ('productos', 'Productos consumidos'),
    ('tiempo', 'Distribución de tiempo'),
    ('eventos', 'Eventos no programados'),
    ('costos', 'Costos'),
])


def _f(v):
    try:
        return float(v) if v not in (None, '') else 0.0
    except (TypeError, ValueError):
        return 0.0


def _principal(reporte):
    checks = list(reporte.mud_checks.all().order_by('check_number'))
    return next((c for c in checks if c.is_primary), checks[0] if checks else None)


def datos_recap(pozo, hasta=None):
    """
    Datos del reporte final. hasta: fecha límite (incluida); None = todos los reportes.
    Devuelve un dict listo para la plantilla y para el Excel.
    """
    from . import volumetria as vol
    from . import control_solidos as cs
    from .volumetria import ErrorVolumetria
    from .control_solidos import ErrorMallas
    from .views_inventario import _linea_de_tiempo, _productos_pozo, _completar_productos, _estado_perdidas
    from .views_control_solidos import _eventos_pozo, _nombres_malla, _contexto_hoyo
    from .models_opcionales import EventoNoProgramado

    qs = pozo.reportes_diarios.order_by('fecha').select_related('intervalo_costo', 'comentarios')
    if hasta:
        qs = qs.filter(fecha__lte=hasta)
    reportes = list(qs)
    header = getattr(pozo, 'well_header_info', None)
    recap = getattr(pozo, 'recap', None)
    moneda = pozo.moneda_simbolo or 'USD'

    base = {
        'pozo': pozo, 'header': header, 'recap': recap, 'moneda': moneda,
        'generado': datetime.now(), 'hasta': hasta, 'reportes': reportes,
        'hay_reportes': bool(reportes), 'avisos': [],
    }
    if not reportes:
        return base
    ultimo = reportes[-1]

    # ---- Simulaciones del pozo, una sola vez ----
    try:
        dias, _, _, ids = _linea_de_tiempo(pozo, ultimo)
        productos = _completar_productos(_productos_pozo(pozo), ids)
        nombres_p = {i: f"{p['descripcion']} ({p['codigo']})" for i, p in productos.items()}
        servicios = {i for i, p in productos.items() if p['servicio']}
    except Exception:
        dias, productos, nombres_p, servicios = [], {}, {}, set()
        base['avisos'].append('No se pudo armar la volumetría del pozo.')
    try:
        eventos_mallas = _eventos_pozo(pozo)
        nombres_m = _nombres_malla(pozo)
    except Exception:
        eventos_mallas, nombres_m = [], {}

    categorias = {c['codigo']: c for c in _estado_perdidas(pozo)['disponibles']}

    filas = []
    perdidas_tot = defaultdict(float)
    vol_tot = defaultdict(float)
    tiempo_tot = OrderedDict()
    costo_acum = 0.0
    anterior = None
    sim_final = None
    dias_con_error = []

    for r in reportes:
        fila = {'reporte': r, 'fecha': r.fecha, 'numero': r.numero_reporte,
                'profundidad': _f(r.profundidad_actual), 'actividad': r.actividad_actual or ''}

        # Hoyo perforado del día (incluye side track y piloto)
        try:
            ctx = _contexto_hoyo(r, anterior)
            fila['avance'] = ctx['avance_ft']
            fila['hoyo_bbl'] = ctx['volumen_hoyo_bbl']
        except Exception:
            fila['avance'] = max(_f(r.profundidad_actual) - (_f(anterior.profundidad_actual) if anterior else 0), 0)
            fila['hoyo_bbl'] = 0.0
        fila['sidetrack'] = bool(_f(getattr(r, 'kickoff_sidetrack_ft', 0)))

        # Lodo (chequeo principal)
        c = _principal(r)
        fila['densidad'] = c.mud_weight if c else None
        fila['embudo'] = c.funnel_viscosity if c else None
        fila['pv'] = c.pv if c else None
        fila['yp'] = c.yp if c else None
        fila['filtrado'] = (c.api_fluid_loss if r.tipo_lodo not in ('OBM', 'SBM') else c.hthp_fluid_loss) if c else None
        cm = getattr(r, 'comentarios', None)
        fila['recap'] = (cm.mud_recap_remarks if cm else '') or ''

        # Volumetría y costos de químicos / ingenieros del día
        chem = ife = 0.0
        sim = None
        if dias:
            try:
                sim = vol.simular(dias, r.id, nombres_p, {}, servicios)
            except ErrorVolumetria:
                dias_con_error.append(r.fecha.strftime('%d/%m/%Y'))
        if sim:
            for cat, v in sim['costo_dia_cat'].items():
                if cat in (1, 2):
                    chem += v
                else:
                    ife += v
            for g in sim['flujos'].values():
                for k in ('aceite', 'agua', 'quimicos', 'recibido', 'devuelto', 'perdida'):
                    vol_tot[k] += g.get(k, 0.0)
            perd_dia = 0.0
            for codigo, por_g in sim['perdidas'].items():
                s = sum(por_g.values())
                perdidas_tot[codigo] += s
                perd_dia += s
            fila['perdidas'] = perd_dia
            real = sim['real_grupo'].get(vol.ACTIVO)
            fila['vol_activo'] = real if real is not None else sim['calc_grupo'].get(vol.ACTIVO)
            sim_final = sim
        else:
            fila['perdidas'] = None
            fila['vol_activo'] = None

        # Equipos y mallas
        equipos = sum(u.costo_diario for u in r.usos_equipo.all())
        mallas = 0.0
        if eventos_mallas:
            try:
                m = cs.simular(eventos_mallas, r.id, nombres_m)
                mallas = sum(x['costo_diario'] for x in m['movimientos'].values())
            except ErrorMallas:
                pass
        fila['costo_quimicos'] = chem
        fila['costo_ingenieria'] = ife
        fila['costo_equipos'] = equipos + mallas
        fila['costo_dia'] = chem + ife + equipos + mallas
        costo_acum += fila['costo_dia']
        fila['costo_acum'] = costo_acum

        for a in r.actividades_tiempo.all().order_by('orden', 'id'):
            tiempo_tot[a.descripcion] = tiempo_tot.get(a.descripcion, 0.0) + _f(a.horas)

        filas.append(fila)
        anterior = r

    if dias_con_error:
        base['avisos'].append('La volumetría no cuadra en: ' + ', '.join(dias_con_error[:5]) +
                              ('…' if len(dias_con_error) > 5 else '') + '. Sus volúmenes y costos de químicos salen en 0.')

    # ---- Intervalos ----
    intervalos = []
    por_intervalo = defaultdict(list)
    for f in filas:
        por_intervalo[f['reporte'].intervalo_costo_id].append(f)
    for iv in pozo.intervalos_revestimiento.all().order_by('numero_intervalo'):
        fs = por_intervalo.get(iv.id, [])
        dens = [_f(f['densidad']) for f in fs if f['densidad'] is not None]
        intervalos.append({
            'iv': iv, 'tipo': iv.get_tipo_display() if iv.tipo else '',
            'dias': len(fs),
            'desde': fs[0]['fecha'] if fs else None, 'hasta': fs[-1]['fecha'] if fs else None,
            'prof_final': fs[-1]['profundidad'] if fs else _f(iv.profundidad_ft),
            'perforado': sum(f['avance'] for f in fs),
            'densidad_max': max(dens) if dens else None,
            'perdidas': sum(f['perdidas'] or 0 for f in fs),
            'costo': sum(f['costo_dia'] for f in fs),
            'costo_planeado': _f(iv.planned_cost) or None,
            'observaciones': iv.observaciones_recomendaciones or '',
            'comentarios': iv.comentarios_recap or '',
        })
    sin_iv = por_intervalo.get(None, [])

    # ---- Pérdidas por categoría ----
    perdidas = []
    for codigo, v in sorted(perdidas_tot.items(), key=lambda x: (x[0] is None, x[0])):
        if v <= 0:
            continue
        cat = categorias.get(codigo, {})
        perdidas.append({'codigo': codigo, 'descripcion': cat.get('descripcion', f"Categoría {codigo}"),
                         'tipo': cat.get('tipo_display', ''), 'volumen': v})
    total_perdidas = sum(p['volumen'] for p in perdidas)

    # ---- Productos consumidos (acumulado del último día) ----
    consumo = []
    if sim_final:
        for pid, s in sim_final['inventario'].items():
            if not (s.get('usado_acum') or s.get('costo_acumulado')):
                continue
            p = productos.get(pid, {})
            consumo.append({'descripcion': p.get('descripcion', pid), 'codigo': p.get('codigo', ''),
                            'unidad': p.get('unidad') or 'servicio', 'tamano': p.get('tamano'),
                            'empaque': p.get('empaque', ''), 'usado': s.get('usado_acum', 0.0),
                            'costo': s.get('costo_acumulado', 0.0)})
        consumo.sort(key=lambda x: -x['costo'])

    # ---- Eventos no programados ----
    eventos = list(EventoNoProgramado.objects.filter(reporte__in=[r.id for r in reportes])
                   .select_related('reporte').order_by('reporte__fecha', 'id'))

    total_tiempo = sum(tiempo_tot.values())
    tiempo = [{'actividad': k, 'horas': v, 'pct': (v / total_tiempo * 100) if total_tiempo else 0}
              for k, v in sorted(tiempo_tot.items(), key=lambda x: -x[1]) if v]

    dens = [_f(f['densidad']) for f in filas if f['densidad'] is not None]
    costos = {
        'quimicos': sum(f['costo_quimicos'] for f in filas),
        'ingenieria': sum(f['costo_ingenieria'] for f in filas),
        'equipos': sum(f['costo_equipos'] for f in filas),
    }
    costos['total'] = sum(costos.values())

    base.update({
        'filas': filas, 'intervalos': intervalos, 'dias_sin_intervalo': len(sin_iv),
        'perdidas': perdidas, 'total_perdidas': total_perdidas,
        'volumenes': dict(vol_tot), 'consumo': consumo, 'eventos': eventos,
        'tiempo': tiempo, 'total_tiempo': total_tiempo, 'costos': costos,
        'kpi': {
            'dias': len(filas),
            'desde': filas[0]['fecha'], 'hasta_fecha': filas[-1]['fecha'],
            'profundidad_final': filas[-1]['profundidad'],
            'perforado': sum(f['avance'] for f in filas),
            'hoyo_bbl': sum(f['hoyo_bbl'] for f in filas),
            'densidad_min': min(dens) if dens else None,
            'densidad_max': max(dens) if dens else None,
            'sidetracks': sum(1 for f in filas if f['sidetrack']),
            'eventos': len(eventos),
            'costo_por_ft': (costos['total'] / sum(f['avance'] for f in filas)) if sum(f['avance'] for f in filas) else None,
        },
    })
    return base


# =====================================================================
# Excel
# =====================================================================

def generar_recap_excel(d, secciones):
    """Libro .xlsx con una hoja por sección elegida. Devuelve (bytes, nombre de archivo)."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    naranja = PatternFill('solid', fgColor='F2994A')
    gris = PatternFill('solid', fgColor='F3F4F6')
    fino = Side(style='thin', color='D1D5DB')
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    blanco_neg = Font(bold=True, color='FFFFFF')
    neg = Font(bold=True)
    envolver = Alignment(wrap_text=True, vertical='top')
    pozo, h, moneda = d['pozo'], d['header'], d['moneda']

    def hoja(titulo, nombre):
        ws = wb.create_sheet(nombre[:31])
        ws['A1'] = f"REPORTE FINAL DEL POZO — {titulo.upper()}"
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f"{pozo.nombre} · {h.operador if h and h.operador else ''} · generado {d['generado']:%d/%m/%Y %H:%M}"
        ws['A2'].font = Font(italic=True, color='6B7280')
        return ws

    def tabla(ws, fila, cabeceras, filas, formatos=None, anchos=None):
        formatos = formatos or {}
        for j, t in enumerate(cabeceras, 1):
            c = ws.cell(row=fila, column=j, value=t)
            c.font, c.fill, c.border = blanco_neg, naranja, borde
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for i, valores in enumerate(filas, 1):
            for j, v in enumerate(valores, 1):
                c = ws.cell(row=fila + i, column=j, value=v)
                c.border = borde
                if j in formatos and isinstance(v, (int, float)):
                    c.number_format = formatos[j]
                if isinstance(v, str) and len(v) > 40:
                    c.alignment = envolver
                if i % 2 == 0:
                    c.fill = gris
        if anchos:
            for j, a in enumerate(anchos, 1):
                ws.column_dimensions[get_column_letter(j)].width = a
        ws.freeze_panes = ws.cell(row=fila + 1, column=1)
        return fila + len(filas) + 2

    ENT, D1, D2 = '#,##0', '#,##0.0', '#,##0.00'
    if not d.get('hay_reportes'):
        ws = hoja('Sin datos', 'Resumen')
        ws['A4'] = 'El pozo todavía no tiene reportes diarios.'
    k = d.get('kpi', {})

    if 'resumen' in secciones and d.get('hay_reportes'):
        ws = hoja('Datos del pozo', 'Resumen')
        datos = [
            ('Pozo', pozo.nombre), ('Operador', h.operador if h else ''), ('Campo / Área', h.field_area if h else ''),
            ('Ubicación', h.ubicacion if h else ''), ('Contratista', h.contratista if h else ''),
            ('Taladro', h.nombre_taladro if h else ''),
            ('Ingenieros', ' / '.join(x for x in ((h.ingeniero_miswaco_1, h.ingeniero_miswaco_2) if h else ()) if x)),
            ('Tipo de lodo principal', h.primary_mud_type_descripcion if h else ''),
            ('Primer reporte', k['desde'].strftime('%d/%m/%Y')), ('Último reporte', k['hasta_fecha'].strftime('%d/%m/%Y')),
            ('Días con reporte', k['dias']), ('Profundidad final (ft)', round(k['profundidad_final'])),
            ('Hoyo perforado (ft)', round(k['perforado'])), ('Volumen de hoyo perforado (bbl)', round(k['hoyo_bbl'])),
            ('Densidad mín. / máx. (lb/gal)', f"{k['densidad_min'] or 0:.1f} / {k['densidad_max'] or 0:.1f}" if k['densidad_max'] else ''),
            ('Side tracks', k['sidetracks']), ('Eventos no programados', k['eventos']),
            ('Pérdidas totales (bbl)', round(d['total_perdidas'])),
            (f"Costo total ({moneda})", round(d['costos']['total'], 2)),
            (f"Costo por pie ({moneda}/ft)", round(k['costo_por_ft'], 2) if k['costo_por_ft'] else ''),
        ]
        for i, (et, v) in enumerate(datos, 4):
            ws.cell(row=i, column=1, value=et).font = neg
            ws.cell(row=i, column=2, value=v)
        ws.column_dimensions['A'].width = 34
        ws.column_dimensions['B'].width = 44

    if 'conclusiones' in secciones and d.get('hay_reportes'):
        ws = hoja('Conclusiones y recomendaciones', 'Conclusiones')
        rc = d['recap']
        fila = 4
        for titulo, texto in (('Resumen del pozo', rc.resumen if rc else ''), ('Conclusiones', rc.conclusiones if rc else ''),
                              ('Recomendaciones', rc.recomendaciones if rc else ''),
                              ('Lecciones aprendidas / buenas prácticas', rc.lecciones if rc else '')):
            ws.cell(row=fila, column=1, value=titulo).font = Font(bold=True, color='F2994A', size=12)
            c = ws.cell(row=fila + 1, column=1, value=texto or '—')
            c.alignment = envolver
            ws.row_dimensions[fila + 1].height = max(30, 15 * (1 + len(texto or '') // 110))
            fila += 3
        ws.column_dimensions['A'].width = 120

    if 'intervalos' in secciones and d.get('hay_reportes'):
        ws = hoja('Resumen por intervalo', 'Intervalos')
        filas = [[iv['iv'].numero_intervalo, iv['tipo'], _f(iv['iv'].casing_od_in) or None, _f(iv['iv'].hole_size_in) or None,
                  iv['desde'].strftime('%d/%m/%Y') if iv['desde'] else '', iv['hasta'].strftime('%d/%m/%Y') if iv['hasta'] else '',
                  iv['dias'], round(iv['prof_final']), round(iv['perforado']),
                  round(iv['densidad_max'], 1) if iv['densidad_max'] else None, round(iv['perdidas']),
                  round(iv['costo'], 2), iv['costo_planeado'], ' '.join(x for x in (iv['observaciones'], iv['comentarios']) if x)]
                 for iv in d['intervalos']]
        tabla(ws, 4, ['N°', 'Tipo', 'OD rev. (in)', 'Hoyo (in)', 'Desde', 'Hasta', 'Días', 'Prof. final (ft)',
                      'Perforado (ft)', 'Densidad máx.', 'Pérdidas (bbl)', f'Costo ({moneda})', f'Costo planeado ({moneda})',
                      'Observaciones y recomendaciones'],
              filas, {12: D2, 13: D2}, [6, 16, 11, 10, 12, 12, 7, 13, 13, 12, 12, 15, 15, 60])

    if 'diario' in secciones and d.get('hay_reportes'):
        ws = hoja('Recap día por día', 'Recap Diario')
        filas = [[f['fecha'].strftime('%d/%m/%Y'), f['numero'], round(f['profundidad']), round(f['avance']),
                  round(_f(f['densidad']), 1) if f['densidad'] is not None else None,
                  round(_f(f['pv'])) if f['pv'] is not None else None, round(_f(f['yp'])) if f['yp'] is not None else None,
                  round(_f(f['filtrado']), 1) if f['filtrado'] is not None else None,
                  round(f['vol_activo']) if f['vol_activo'] is not None else None,
                  round(f['perdidas']) if f['perdidas'] is not None else None,
                  round(f['costo_dia'], 2), round(f['costo_acum'], 2),
                  ('[SIDE TRACK] ' if f['sidetrack'] else '') + (f['actividad'] or ''), f['recap']]
                 for f in d['filas']]
        tabla(ws, 4, ['Fecha', 'N°', 'Prof. (ft)', 'Avance (ft)', 'Densidad (lb/gal)', 'PV', 'YP', 'Filtrado',
                      'Vol. activo (bbl)', 'Pérdidas (bbl)', f'Costo día ({moneda})', f'Acumulado ({moneda})',
                      'Actividad', 'Resumen del día'],
              filas, {11: D2, 12: D2}, [11, 5, 10, 10, 10, 6, 6, 8, 11, 10, 13, 14, 28, 70])

    if 'volumenes' in secciones and d.get('hay_reportes'):
        ws = hoja('Volúmenes y pérdidas', 'Volumenes')
        v = d['volumenes']
        fila = tabla(ws, 4, ['Concepto', 'Volumen (bbl)'], [
            ['Fluido base agregado', round(v.get('aceite', 0))], ['Agua agregada', round(v.get('agua', 0))],
            ['Volumen de químicos', round(v.get('quimicos', 0))], ['Lodo entero recibido', round(v.get('recibido', 0))],
            ['Devuelto', round(v.get('devuelto', 0))], ['Perdido y descartado', round(v.get('perdida', 0))],
        ], None, [36, 16])
        ws.cell(row=fila, column=1, value='Pérdidas por categoría').font = neg
        tabla(ws, fila + 1, ['Código', 'Categoría', 'Tipo', 'Volumen (bbl)', '% del total'],
              [[p['codigo'], p['descripcion'], p['tipo'], round(p['volumen']),
                round(p['volumen'] / d['total_perdidas'] * 100, 1) if d['total_perdidas'] else 0] for p in d['perdidas']]
              + [['', 'TOTAL', '', round(d['total_perdidas']), 100 if d['total_perdidas'] else 0]],
              None, [36, 36, 14, 14, 12])

    if 'productos' in secciones and d.get('hay_reportes'):
        ws = hoja('Productos consumidos', 'Productos')
        filas = [[p['codigo'], p['descripcion'], f"{p['tamano']:g} {p['unidad']} {p['empaque'] or ''}".strip()
                  if p['tamano'] else p['unidad'], round(p['usado'], 2), round(p['costo'], 2)] for p in d['consumo']]
        filas.append(['', 'TOTAL', '', None, round(sum(p['costo'] for p in d['consumo']), 2)])
        tabla(ws, 4, ['Código', 'Producto', 'Tamaño', 'Cantidad usada', f'Costo ({moneda})'], filas,
              {4: D2, 5: D2}, [14, 40, 22, 15, 16])

    if 'tiempo' in secciones and d.get('hay_reportes'):
        ws = hoja('Distribución de tiempo', 'Tiempo')
        filas = [[t['actividad'], round(t['horas'], 2), round(t['pct'], 1)] for t in d['tiempo']]
        filas.append(['TOTAL', round(d['total_tiempo'], 2), 100 if d['total_tiempo'] else 0])
        tabla(ws, 4, ['Actividad', 'Horas', '% del tiempo'], filas, {2: D2}, [40, 12, 12])

    if 'eventos' in secciones and d.get('hay_reportes'):
        ws = hoja('Eventos no programados', 'Eventos')
        filas = [[e.reporte.fecha.strftime('%d/%m/%Y'), e.get_categoria_display(), e.tipo_problema, e.descripcion,
                  e.causa, e.horas_perdidas, round(e.volumen_perdido_bbl) if e.volumen_perdido_bbl is not None else None,
                  e.costo] for e in d['eventos']]
        if not filas:
            filas = [['', 'Sin eventos registrados', '', '', '', None, None, None]]
        tabla(ws, 4, ['Fecha', 'Categoría', 'Problema', 'Descripción', 'Causa', 'Tiempo perdido (h)',
                      'Volumen perdido (bbl)', f'Costo ({moneda})'], filas, {8: D2}, [11, 20, 24, 45, 35, 12, 12, 14])

    if 'costos' in secciones and d.get('hay_reportes'):
        ws = hoja('Costos', 'Costos')
        c = d['costos']
        fila = tabla(ws, 4, ['Concepto', f'Costo ({moneda})'], [
            ['Químicos / Personal DF', round(c['quimicos'], 2)], ['Ingeniero IFE / Control de sólidos', round(c['ingenieria'], 2)],
            ['Equipos y mallas', round(c['equipos'], 2)], ['TOTAL', round(c['total'], 2)]], {2: D2}, [38, 18])
        ws.cell(row=fila, column=1, value='Por intervalo').font = neg
        tabla(ws, fila + 1, ['Intervalo', 'Días', f'Costo real ({moneda})', f'Costo planeado ({moneda})', 'Diferencia'],
              [[f"{iv['iv'].numero_intervalo} {iv['tipo']}".strip(), iv['dias'], round(iv['costo'], 2), iv['costo_planeado'],
                round(iv['costo'] - iv['costo_planeado'], 2) if iv['costo_planeado'] else None] for iv in d['intervalos']],
              {3: D2, 4: D2, 5: D2}, [38, 8, 18, 18, 14])

    if not wb.worksheets:
        ws = wb.create_sheet('Resumen')
        ws['A1'] = 'No se eligió ninguna sección.'
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    salida = io.BytesIO()
    wb.save(salida)
    nombre = f"Reporte_Final_{pozo.nombre}_{d['generado']:%Y%m%d}.xlsx".replace(' ', '_')
    return salida.getvalue(), nombre
