"""
Crea operaciones/plantillas/reporte_diario_onetrax.xlsx a partir del Excel original de ONE-TRAX
(Mud Report 16 PERLA-1X.xlsx): conserva formato, combinaciones, anchos y áreas de impresión,
traduce los rótulos al español y vacía todos los datos y fórmulas externas.
Se corre una sola vez (fuera de Django); el generador trabaja sobre la plantilla resultante.
"""
import sys
import openpyxl

ORIGEN = sys.argv[1] if len(sys.argv) > 1 else 'plantilla_origen.xlsx'
DESTINO = sys.argv[2] if len(sys.argv) > 2 else 'reporte_diario_onetrax.xlsx'

HOJAS = {
    'WBM Check': 'Lodo Base Agua',
    'CALDRIL Check': 'Lodo CALDRIL',
    'OBM Check': 'Lodo Base Aceite',
    'Synthetic-Based Mud': 'Lodo Base Sintetica',
    'Extra WBM Labels': 'Prop Extra Base Agua',
    'Extra OBM-SBM Labels': 'Prop Extra Aceite-Sint',
    'Vol. Info': 'Contabilidad de Volumen',
    'Chem. Inv (DF)': 'Inv Quimico (DF)',
    'Chem. Inv (DF-Order by Name)': 'Inv Quimico (por nombre)',
    'Chem. Inv (Full)': 'Inv Quimico (completo)',
    'Equipment': 'Equipos',
    'Screen Inventory': 'Inventario de Mallas',
}

T = {
    # Encabezados comunes
    'Operator : ': 'Operador : ', 'Operator :': 'Operador :', 'Report For : ': 'Reporte para : ',
    'Well Name : ': 'Pozo : ', ' Well Name :': ' Pozo :', 'Well Name:': 'Pozo :',
    'Contractor: ': 'Contratista : ', 'Field/Area :': 'Campo/Área :', 'Description :': 'Descripción :',
    'Location :': 'Ubicación :', 'Location:': 'Ubicación :', 'Water Depth :': 'Prof. de Agua :',
    'Rig Name :': 'Taladro :', 'Depth/TVD :': 'Prof./TVD :', 'Date :': 'Fecha :', 'Date : ': 'Fecha : ',
    'Spud Date :': 'Fecha Spud :', 'Mud Type :': 'Tipo de Lodo :', 'Activity :': 'Actividad :',
    'Report No:': 'Reporte N° :', 'Report No.:': 'Reporte N° :', 'Page ': 'Página ',
    # Reporte de lodo
    'DRILLING ASSEMBLY': 'SARTA DE PERFORACIÓN', 'CASING': 'REVESTIDOR', 'CASING (*TVD)': 'REVESTIDOR (*TVD)',
    'MUD VOLUME': 'VOLUMEN DE LODO', 'CIRCULATION DATA': 'DATOS DE CIRCULACIÓN', 'Hole': 'Hoyo',
    'Active Pits': 'Fosas Activas', 'Pump Make': 'Marca de Bomba', 'Pump Liner x Stk': 'Camisa x Carrera',
    'Total Circulating Volume': 'Volumen Total en Circulación', 'Pump Capacity gal/stk': 'Capacidad gal/emb', 'Pump stk/min': 'Bomba emb/min',
    'Depth Drilled Last 24hr': 'Prof. Perforada Últ. 24 h', 'Depth Drilled Last 24 hr': 'Prof. Perforada Últ. 24 h',
    ' Volume Drilled Last 24hr': ' Volumen Perforado Últ. 24 h', ' Volume Drilled Last 24 hr': ' Volumen Perforado Últ. 24 h',
    'Flow Rate': 'Caudal', 'Pump Pressure': 'Presión de Bomba', 'Bottoms Up': 'Fondo Arriba',
    'Total Circulation': 'Circulación Total', 'MUD PROPERTIES': 'PROPIEDADES DEL LODO',
    'PRODUCTS USED Last 24 hr': 'PRODUCTOS USADOS Últ. 24 h', 'Products ': 'Productos ', 'Size': 'Tamaño',
    'Amount': 'Cantidad', 'Sample From': 'Muestra de', 'FlowLine Temp': 'Temp. L. Flujo',
    'Depth/TVD': 'Prof./TVD', 'Mud Weight /Temp': 'Densidad /Temp', 'Mud Weight': 'Densidad',
    'Funnel Viscosity': 'Viscosidad Embudo', 'Rheology Temp': 'Temp. Reología', 'R600/R300': 'L600/L300',
    'R200/R100': 'L200/L100', 'R6/R3': 'L6/L3', 'PV': 'VP', 'YP': 'PC', '10s/10m/30m Gel': 'Geles 10s/10m/30m',
    'API Fluid Loss': 'Filtrado API', 'HTHP Fluid Loss': 'Filtrado HPHT', 'Cake APT/HT': 'Revoque API/HT',
    'Solids': 'Sólidos', 'Oil/Water': 'Aceite/Agua', 'Sand': 'Arena', 'MBT': 'MBT', 'pH / Temp': 'pH / Temp',
    'Alkal Mud (Pm)': 'Alcal. Lodo (Pm)', 'Pf/Mf': 'Pf/Mf', 'Chlorides': 'Cloruros', 'Hardness (Ca++)': 'Dureza (Ca++)',
    'Unc Ret Solids': 'Sólidos s/Corr.', 'Correct Solids': 'Sólidos Corregidos', 'Oil': 'Aceite',
    'Synthetic': 'Sintético', 'Uncorr Water': 'Agua s/Corr.', 'Oil/Water Ratio': 'Rel. Aceite/Agua',
    'Synthetic/Water Ratio': 'Rel. Sint./Agua', 'Alkal Mud (Pom)': 'Alcal. Lodo (Pom)',
    'Alkal Mud (Psm)': 'Alcal. Lodo (Psm)', 'Cl- Whole Mud': 'Cl- Lodo Entero', 'Salt': 'Sal', 'Lime': 'Cal',
    'Emul Stability': 'Estab. Eléctrica',
    'SOLIDS CONTROL EQUIPMENT Last 24 hr': 'EQUIPOS DE CONTROL DE SÓLIDOS Últ. 24 h', 'Type': 'Tipo',
    'Model/Size': 'Modelo/Mallas', 'Hrs Used': 'Horas Uso', 'MUD PROPERTY SPECS': 'ESPECIFICACIÓN DEL LODO',
    'Actual': 'Real', 'Weight': 'Densidad', 'Viscosity': 'Viscosidad', 'Filtrate': 'Filtrado',
    'Reserve Volume': 'Vol. Reserva', 'REMARKS AND TREATMENT': 'OBSERVACIONES Y TRATAMIENTO',
    'REMARKS': 'OBSERVACIONES', 'TIME DISTRIBUTION Last 24 hrs': 'DIST. DE TIEMPO Últ. 24 h',
    'SOLIDS ANALYSIS (%/lb/bbl)': 'ANÁLISIS DE SÓLIDOS (%/lb/bbl)', 'SOLIDS ANALYSIS': 'ANÁLISIS DE SÓLIDOS',
    'RHEOLOGY & HYDRAULICS': 'REOLOGÍA E HIDRÁULICA', 'Oil Added': 'Aceite Agregado',
    'Synthetic Added': 'Sintético Agregado', 'Water Added': 'Agua Agregada', 'Mud Received': 'Lodo Recibido',
    'Mud Returned': 'Lodo Devuelto', 'NaCl': 'NaCl', 'KCl': 'KCl', 'Low Gravity': 'Baja Gravedad',
    'Bentonite': 'Bentonita', 'Drill Solids': 'Sólidos Perforados', 'Weight Material': 'Material Densificante',
    'Chemical Conc': 'Conc. Química', 'Inert/React': 'Inerte/Reactivo', 'Average SG': 'GE Promedio',
    'Salt Wt%': 'Sal % peso', 'Salt Conc': 'Conc. de Sal', 'Adjusted Solids': 'Sólidos Ajustados',
    'Average SG Solids': 'GE Promedio Sólidos', 'Low Gravity %': 'Baja Gravedad %',
    'Low Gravity Wt.': 'Baja Gravedad Peso', 'High Gravity %': 'Alta Gravedad %', 'High Gravity Wt.': 'Alta Gravedad Peso',
    'Calcium Chloride (wt %)': 'Cloruro de Calcio (% peso)', 'Brine Specific Gravity': 'GE de la Salmuera',
    'Corrected Solids (%)': 'Sólidos Corregidos (%)', 'Average SG of Solids': 'GE Promedio de Sólidos',
    'Low Gravity Solids (%)': 'Sólidos Baja Gravedad (%)', 'High Gravity Solids (%)': 'Sólidos Alta Gravedad (%)',
    'np/na': 'np/na', 'Kp/Ka': 'Kp/Ka', 'Bit Pressure Loss %': 'Pérd. Mecha psi/%',
    'Bit Pressure Loss/%': 'Pérd. Mecha psi/%', 'Bit HHP/HSI': 'HHP/HSI Mecha', 'Jet Velocity': 'Velocidad de Chorro',
    'Va Pipe': 'Va Tubería', 'Va Collars': 'Va Portamechas', 'Cva Pipe': 'Vc Tubería', 'Cva Collars': 'Vc Portamechas',
    'ECD at Shoe': 'DEC en Zapata', 'ECD at TD': 'DEC en Fondo', 'DAILY COST': 'COSTO DIARIO',
    'CUMULATIVE COST': 'COSTO ACUMULADO', 'Drlg Fluids': 'Fluidos Perf.', 'S.R.E. & Engr.': 'C.S. e Ing.',
    'Screens': 'Equipos', 'TOTALS': 'TOTALES', 'M-I ENGR / PHONE': 'INGENIERO / TELÉFONO',
    'RIG PHONE': 'TEL. TALADRO', 'WAREHOUSE PHONE': 'TEL. ALMACÉN', 'EXTRA  PROPERTIES': 'PROPIEDADES EXTRA',
    # Contabilidad de volumen
    'MUD VOLUME ACCOUNTING': 'CONTABILIDAD DE VOLUMEN', 'Daily Report': 'Reporte Diario',
    'TANK': 'FOSA', 'CAPACITY': 'CAPACIDAD', 'WEIGHT': 'DENSIDAD', 'VOLUME': 'VOLUMEN', 'CLASS': 'TIPO',
    'SUM PIT VOLUMES': 'SUMA DE FOSAS', 'Active': 'Activa', 'Reserve': 'Reserva', 'Premix': 'Premezcla',
    'OTHER VOLUME': 'OTROS VOLÚMENES', 'MUD IN HOLE (bbls)': 'LODO EN EL HOYO (bbl)', 'MUD IN HOLE:': 'LODO EN EL HOYO:',
    'ANNULUS': 'ANULAR', 'PIPE': 'SARTA', 'BELOW BIT': 'BAJO LA MECHA', 'TOTAL': 'TOTAL',
    'Total Hole Volume': 'Vol. Total del Hoyo', 'Volume Not Mud': 'Vol. que no es Lodo', 'Mud Volume': 'Volumen de Lodo',
    'VOLUME BALANCE (bbls)': 'BALANCE DE VOLUMEN (bbl)', 'LOSS BREAKDOWN (bbls)': 'DESGLOSE DE PÉRDIDAS (bbl)',
    'IN/TO:': 'EN/HACIA:', 'ACTIVE': 'ACTIVO', 'RESERVE': 'RESERVA', 'PREMIX': 'PREMEZCLA',
    'Start Volume': 'Volumen Inicial', 'Vol Chem Added': 'Vol. Químicos Agregados', 'Chem Not For Mud': 'Químicos no para Lodo',
    'Total Volume Built': 'Vol. Total Construido', 'Received': 'Recibido', 'Return': 'Devuelto',
    'From Active To': 'Del Activo Hacia', 'From Reserve To': 'De Reserva Hacia', 'From Premix To': 'De Premezcla Hacia',
    'Daily Loss': 'Pérdida del Día', 'Final Volume': 'Volumen Final', 'Total Loss': 'Pérdida Total',
    # Inventario químico
    'WELLSITE CHEMICAL INVENTORY': 'INVENTARIO QUÍMICO EN LOCACIÓN', 'Cost Summary': 'Resumen de Costos',
    'Total Daily Cost:': 'Costo Total Diario:', 'Total Daily Tax:': 'Impuesto Diario:', 'Cumulative Cost:': 'Costo Acumulado:',
    'Product': 'Producto', 'Unit ': 'Tamaño ', 'Unit': 'Precio', 'Start': 'Inicial', 'Daily': 'Diario',
    'Daily ': 'Costo', 'Cum': 'Acum.', 'Cum.': 'Acum.', 'Final': 'Final', 'Amt.': 'Cant.', 'Used': 'Usado',
    "Rec'd": 'Recibido', 'Stock': 'Existencia', 'Cost': 'Costo', 'Price': 'Unitario',
    # Equipos
    'Unit Cost': 'Costo Unitario', '(F)ull or': '(C)ompleto o', '(S)tand-by': '(S)tand-by', 'Code': 'Código',
    'Daily Rental': 'Renta Diaria', 'Equipment Description': 'Descripción del Equipo', 'Full': 'Completo',
    'Stand-by': 'Stand-by', 'Today': 'Hoy', 'Drilling  Fluids Equipment': 'Equipos de Fluidos de Perforación',
    'Other Solids Removal Equipment': 'Otros Equipos de Remoción de Sólidos', 'Additional Equipment': 'Equipos Adicionales',
    'Total': 'Total',
    # Mallas
    'Screen Name': 'Malla', 'Mesh': 'Mesh', 'Installed Today': 'Instaladas Hoy', 'New': 'Nuevas', 'for Well': 'en el Pozo',
    'Daily Cost:': 'Costo Diario:',
}

# Rótulos que en algunas hojas cambian de sentido (inventario / mallas) y no deben tomarse del diccionario general.
POR_HOJA = {
    'Inventario de Mallas': {'Unit': 'Costo', 'Used': 'Usadas', 'Daily': 'Costo', 'Size': 'Tamaño',
                             'Cost': 'Unitario', 'Start': 'Inicial', 'Final': 'Final'},
    'Equipos': {'Used': 'Usado', 'Unit Cost': 'Costo Unitario'},
}
# Para las hojas de inventario químico las columnas son de dos filas: fila 9 y fila 10.
INV_FILA_9 = {'Product': 'Producto', 'Unit ': 'Tamaño', 'Unit': 'Precio', 'Start': 'Inicial', 'Daily': 'Diario',
              'Cum': 'Acum.', 'Cum.': 'Acum.', 'Final': 'Final', 'Daily ': 'Costo', 'Unit  ': 'Precio'}
INV_FILA_10 = {'Size': 'Unidad', 'Price': 'Unitario', 'Amt.': 'Cant.', 'Used': 'Usado', "Rec'd": 'Recibido',
               'Return': 'Devuelto', 'Stock': 'Existencia', 'Cost': 'Costo'}

TITULOS = {  # celdas de título que eran fórmulas
    'Lodo Base Agua': {'J2': 'REPORTE DE LODO BASE AGUA N° '},
    'Lodo CALDRIL': {'J2': 'REPORTE DE LODO CALDRIL N° '},
    'Lodo Base Aceite': {'J2': 'REPORTE DE LODO BASE ACEITE N° '},
    'Lodo Base Sintetica': {'J2': 'REPORTE DE LODO BASE SINTÉTICA N° '},
    'Prop Extra Base Agua': {'J2': 'PROPIEDADES EXTRA BASE AGUA N° '},
    'Prop Extra Aceite-Sint': {'J2': 'PROPIEDADES EXTRA ACEITE/SINTÉTICO N° '},
    'Equipos': {'I2': 'Uso y Costo de Equipos'},
    'Inventario de Mallas': {'I1': 'Inventario y Uso de Mallas'},
}


def traducir(hoja, celda, texto):
    fila = celda.row
    if hoja.startswith('Inv Quimico'):
        base = (fila - 1) % 60 + 1   # las páginas se repiten cada 60 filas
        if base == 9 and texto in INV_FILA_9:
            return INV_FILA_9[texto]
        if base == 10 and texto in INV_FILA_10:
            return INV_FILA_10[texto]
    especial = POR_HOJA.get(hoja, {})
    if texto in especial:
        return especial[texto]
    return T.get(texto)


def main():
    wb = openpyxl.load_workbook(ORIGEN)
    for ws in list(wb.worksheets):
        if ws.title not in HOJAS:
            wb.remove(ws)
    wb._external_links = []
    # Nombres definidos que apuntan a hojas borradas o a libros externos
    for nombre in list(wb.defined_names.keys()) if hasattr(wb.defined_names, 'keys') else []:
        del wb.defined_names[nombre]

    for ws in wb.worksheets:
        viejo = ws.title
        ws.title = HOJAS[viejo]
        ws.sheet_state = 'visible'
        ws._images = []        # logos de M-I SWACO: la plantilla va sin logo
        for nombre in list(ws.defined_names.keys()) if hasattr(ws, 'defined_names') else []:
            if nombre != '_xlnm.Print_Area':
                del ws.defined_names[nombre]
        ws.data_validations.dataValidation = []
        ws.conditional_formatting = type(ws.conditional_formatting)()
        sin_traducir = set()
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None or type(c).__name__ == 'MergedCell':
                    continue
                if isinstance(v, str) and not v.strip().startswith('='):
                    nuevo = traducir(ws.title, c, v)
                    if nuevo is not None:
                        c.value = nuevo
                        continue
                    sin_traducir.add(v)
                c.value = None
        if ws.title.startswith('Inv Quimico'):
            from openpyxl.styles import PatternFill, Border
            for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=14, max_col=ws.max_column):
                for c in fila:
                    if type(c).__name__ != 'MergedCell':
                        c.value = None
                        c.fill = PatternFill()
                        c.border = Border()
            for rango in [str(r) for r in ws.merged_cells.ranges]:
                if rango.startswith(('Q', 'R', 'S')):
                    ws.unmerge_cells(rango)
            ws.print_area = 'A1:M180'
        for coord, texto in TITULOS.get(ws.title, {}).items():
            ws[coord].value = texto
        # En la hoja base agua (guardada como valores) los rótulos de análisis de sólidos venían con datos al lado.
        if sin_traducir:
            print(ws.title, 'vaciado (datos):', sorted(sin_traducir)[:12], '...' if len(sin_traducir) > 12 else '')

    wb.active = 0
    wb.save(DESTINO)
    print('Plantilla creada:', DESTINO)


if __name__ == '__main__':
    main()
